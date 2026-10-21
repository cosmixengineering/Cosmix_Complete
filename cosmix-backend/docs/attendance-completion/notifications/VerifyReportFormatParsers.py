"""Independent readers for synthetic output from VerifyReportFormats.cs; no DB/private storage.
Run with the bundled Python and an explicit synthetic artifact folder. Poppler PNG review is separate.
"""
import csv
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import pdfplumber
from pypdf import PdfReader

folder = Path(sys.argv[1]).resolve()
expected = json.loads((folder / "expected.json").read_text(encoding="utf-8"))
rows = expected["Rows"]
metadata = expected["Metadata"]
checks = 0


def check(condition, label):
    global checks
    if not condition:
        raise AssertionError(label)
    checks += 1


with (folder / "attendance-maximum.csv").open(encoding="utf-8-sig", newline="") as source:
    csv_rows = list(csv.reader(source))
header_index = next(i for i, row in enumerate(csv_rows) if row and row[0] == "EmployeeId")
check(len(csv_rows[header_index + 1 :]) == 775, "independent CSV parser preserves all rows")
csv_metadata = dict(row for row in csv_rows[:header_index] if len(row) == 2)
check(csv_metadata["Generated at UTC"].startswith("2026-10-09T12:34:56"), "CSV generation metadata")
check(csv_metadata["Lock status at generation"] == "mixed" and "2026-10-31" in csv_metadata["Policy period"], "CSV lock/policy period metadata")
check(csv_metadata["Employee filter 25"] == metadata["EmployeeIds"][-1], "CSV complete employee filters")
check(csv_rows[-1][11] == rows[-1]["InputFingerprint"], "CSV full source fingerprint")

xlsx = folder / "attendance-maximum.xlsx"
with zipfile.ZipFile(xlsx) as archive:
    for item in archive.namelist():
        if item.endswith((".xml", ".rels")):
            ElementTree.fromstring(archive.read(item))
    check(archive.testzip() is None, "OOXML ZIP integrity and independent XML parsing")
book = openpyxl.load_workbook(xlsx, read_only=False, data_only=False)
check(book.sheetnames == ["Report metadata", "Attendance"], "openpyxl reads real workbook and sheets")
sheet = book["Attendance"]
check(sheet.max_row == 776 and sheet.max_column == 13, "workbook dimensions")
check(sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:M776", "workbook usable frozen headers and filter")
quantities = {4: "WorkedSeconds", 5: "WorkedMinutes", 6: "OpenSessions", 7: "SessionCount", 8: "ExceptionCount"}
for index, source in enumerate(rows, start=2):
    for column, key in quantities.items():
        if sheet.cell(index, column).data_type != "n" or sheet.cell(index, column).value != source[key]:
            raise AssertionError(f"numeric cell {index}/{column}")
    if sheet.cell(index, 1).value != source["EmployeeId"] or sheet.cell(index, 12).value != source["InputFingerprint"]:
        raise AssertionError(f"identity/fingerprint row {index}")
    if source["SummaryRevision"] is not None and (sheet.cell(index, 11).data_type != "n" or sheet.cell(index, 11).value != source["SummaryRevision"]):
        raise AssertionError(f"numeric revision {index}")
check(True, "all 775 rows retain full IDs/hashes and numeric factual quantities/revisions")
check(not any(cell.data_type == "f" for cells in sheet for cell in cells), "no executable workbook formulas")
xlsx_metadata = dict(book["Report metadata"].iter_rows(min_row=2, values_only=True))
check(xlsx_metadata["Report ID"] == metadata["ReportId"] and xlsx_metadata["Employee filter 25"] == metadata["EmployeeIds"][-1], "workbook generation provenance and employee filters")
check(xlsx_metadata["Policy status"] == "factual_only_no_payroll_calculation" and xlsx_metadata["Lock status at generation"] == "mixed", "workbook factual policy and lock metadata")

pdf = folder / "attendance-maximum.pdf"
reader = PdfReader(pdf, strict=True)
check(len(reader.pages) == 76 and not reader.is_encrypted, "independent strict PDF parser and complete pagination")
text = "\n".join(page.extract_text() for page in reader.pages)
check(text.count("Status: ") == 775 and text.count("Input: ") == 775, "PDF all rows and source fingerprints retained")
check(all(employee in text for employee in metadata["EmployeeIds"]), "PDF all employee filters and identifiers retained")
check(metadata["ReportId"] in text and "2026-10-09T12:34:56" in text and "Policy calculations are unavailable" in text and "mixed" in text, "PDF visible generation/filter/policy/lock metadata")
check(rows[-1]["InputFingerprint"] in text and "Page 76 of 76" in text, "PDF complete final page and fingerprint")
with pdfplumber.open(pdf) as document:
    for number, page in enumerate(document.pages, start=1):
        for char in page.chars:
            if char["x0"] < 35 or char["x1"] > 558 or char["top"] < 25 or char["bottom"] > 814:
                raise AssertionError(f"PDF text outside intended page area on page {number}: {char}")
    check(True, "independent PDF geometry: all glyphs stay within intended page area")
print(f"Independent report format parsers: {checks} checks passed; rendered PNG visual inspection remains separate.")
