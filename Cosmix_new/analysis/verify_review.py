import json
from pathlib import Path
import pypdfium2 as pdfium
from pypdf import PdfReader
from PIL import Image, ImageDraw
from docx import Document

ROOT = Path(__file__).resolve().parent
QA = ROOT / 'review-qa'
pdf = PdfReader(QA / 'review.pdf')
texts = [page.extract_text() or '' for page in pdf.pages]
short = [{'page': i + 1, 'characters': len(text.strip()), 'start': text[:100]} for i, text in enumerate(texts) if len(text.strip()) < 400]
doc = Document(ROOT / 'Cosmix_New_Code_Review_and_SRS_Gap_Assessment.docx')
headings = [p.text for p in doc.paragraphs if p.style.name == 'Heading 3']
assert len(headings) == 40, len(headings)
assert all(len(text.strip()) > 100 for text in texts), short
assert all('Cosmix New code review' in text for text in texts)
text = '\n'.join(texts)
for expected in ['2968', '174', '140', '315', '277', '342', '257', '318', 'Customer application', 'Office employee mobile panel', 'Face', 'OTP']:
    assert expected.casefold() in text.casefold(), expected
document = pdfium.PdfDocument(QA / 'review.pdf')
page_paths = []
for i in range(len(document)):
    image = document[i].render(scale=1.65).to_pil()
    path = QA / f'page-{i+1:03d}.png'
    image.save(path)
    page_paths.append(path)
for batch in range(0, len(page_paths), 6):
    canvas = Image.new('RGB', (1530, 1430), 'white')
    draw = ImageDraw.Draw(canvas)
    for n, path in enumerate(page_paths[batch:batch+6]):
        img = Image.open(path)
        img.thumbnail((495, 670))
        x, y = (n % 3) * 510, (n // 3) * 715
        draw.text((x + 5, y + 5), f'Page {batch+n+1}', fill='black')
        canvas.paste(img, (x + 5, y + 25))
    canvas.save(QA / f'contact-{batch//6+1:02d}.png')
summary = {'pages': len(texts), 'findings': len(headings), 'shortPages': short, 'contactSheets': (len(texts)+5)//6}
(ROOT / 'render_verification.json').write_text(json.dumps(summary, indent=2), encoding='utf8')
print(json.dumps(summary, indent=2))
