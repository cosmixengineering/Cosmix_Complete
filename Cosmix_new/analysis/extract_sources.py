import json
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parent
SOURCES = {
    'development_report': ROOT.parent / 'HMS_ERP_FSM_Comprehensive_Development_Report.docx',
    'original_srs': Path('C:/Users/Huma.Khan/Desktop/SOFTWARE REQUIREMENTS SPECIFICATION.docx'),
    'detailed_srs': Path('C:/Users/Huma.Khan/Desktop/Cosmix_ERP_FSM_Detailed_SRS_v0.4.docx'),
    'hr_srs': Path('C:/Users/Huma.Khan/Downloads/HR_Attendance_Employee_Mobile_SRS_v1.2.docx'),
}

for name, path in SOURCES.items():
    doc = Document(path)
    data = {
        'source': str(path),
        'paragraphs': [{'style': p.style.name, 'text': p.text} for p in doc.paragraphs if p.text.strip()],
        'tables': [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables],
    }
    (ROOT / f'{name}.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf8')
    text = '\n'.join(p['text'] for p in data['paragraphs'])
    for index, table in enumerate(data['tables'], 1):
        text += f'\n\nTABLE {index}\n' + '\n'.join(' | '.join(row) for row in table)
    (ROOT / f'{name}.txt').write_text(text, encoding='utf8')
    print(f'{name}: {len(data["paragraphs"])} paragraphs, {len(data["tables"])} tables')
