from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Huma.Khan\Downloads\Cosmix_Erp_software")
SOURCE = ROOT / "docs/superpowers/plans/2026-10-12-cosmix-module-integration-review-plan.md"
OUTPUT = ROOT / "Cosmix_new/analysis/Cosmix_Module_Integration_and_Delivery_Plan_2026_10_12.docx"


def element(name, **attrs):
    node = OxmlElement(name)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    return node


def inline(paragraph, text):
    for token in re.split(r"(\*\*[^*]+\*\*)", text):
        if not token:
            continue
        run = paragraph.add_run(token[2:-2] if token.startswith("**") else token)
        run.bold = token.startswith("**")


def table(document, rows):
    widths = [1.64, 1.74, 3.52]
    result = document.add_table(rows=1, cols=3)
    result.autofit = False
    result.alignment = 1
    props = result._tbl.tblPr
    borders = element("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element(f"w:{edge}", **{"w:val": "single", "w:sz": 5, "w:color": "D9D9D9"}))
    props.append(borders)
    margins = element("w:tblCellMar")
    for side, amount in (("top", 95), ("bottom", 95), ("left", 120), ("right", 120)):
        margins.append(element(f"w:{side}", **{"w:w": amount, "w:type": "dxa"}))
    props.append(margins)
    for index, values in enumerate(rows):
        row = result.rows[0] if index == 0 else result.add_row()
        row._tr.get_or_add_trPr().append(element("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(element("w:tblHeader"))
        for col, value in enumerate(values):
            cell = row.cells[col]
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().append(element("w:shd", **{"w:fill": "353535" if index == 0 else ("F1F5F8" if index % 2 == 0 else "FFFFFF")}))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if index == 0 else "000000")
                if index == 0:
                    run.bold = True
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(3)


document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
    style = document.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.12
    style.paragraph_format.space_after = Pt(6)
    props = style.element.find(qn("w:pPr"))
    if props is not None:
        for border in list(props.findall(qn("w:pBdr"))):
            props.remove(border)
    fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is not None:
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attr}"), None)

document.styles["Title"].font.size = Pt(25)
document.styles["Title"].paragraph_format.space_after = Pt(12)
for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11.5)):
    document.styles[name].font.size = Pt(size)
    document.styles[name].font.bold = True
    document.styles[name].paragraph_format.space_before = Pt(10)
    document.styles[name].paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = "COSMIX    MODULE INTEGRATION PLAN    REVIEW DRAFT"
header.runs[0].font.size = Pt(8)
header.paragraph_format.space_after = Pt(0)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Review draft  |  12 October 2026  |  Page ").font.size = Pt(8)
field = element("w:fldSimple", **{"w:instr": "PAGE"})
footer._p.append(field)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
index = 0
while index < len(lines):
    line = lines[index].strip()
    index += 1
    if not line or line.startswith("> For agentic workers:"):
        continue
    if line.startswith(("**Architecture:**", "**Tech Stack:**", "**Spec:**")):
        continue
    if line == "<!-- page -->":
        next_heading = lines[index].strip() if index < len(lines) else ""
        starts = ("## 2 Confirmed", "## 6 Inventory Integration", "## 9 Complete SRS Module Roadmap",
                  "## 11 Security and Integrity", "## 15 Implementation Packages and Verification",
                  "## 19 Source Map")
        if next_heading.startswith(starts) and not next_heading.endswith(("Continued", "Final Groups")):
            document.add_page_break()
        continue
    if line.startswith("|"):
        rows = [[cell.strip() for cell in line.strip("|").split("|")]]
        while index < len(lines) and lines[index].strip().startswith("|"):
            row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
            index += 1
            if all(re.fullmatch(r":?-+:?", cell) for cell in row):
                continue
            rows.append(row)
        table(document, rows)
        continue
    if line.startswith("# "):
        document.add_paragraph(line[2:], "Title")
    elif line.startswith("## "):
        heading = document.add_paragraph(line[3:], "Heading 1")
        if line.startswith("## 18 Your Review"):
            heading.paragraph_format.page_break_before = True
    elif line.startswith("### "):
        document.add_paragraph(line[4:], "Heading 2")
    elif line.startswith("- [ ] "):
        paragraph = document.add_paragraph(style="List Bullet")
        inline(paragraph, "Review item  " + line[6:])
    elif line.startswith("- "):
        paragraph = document.add_paragraph(style="List Bullet")
        inline(paragraph, line[2:])
    elif re.match(r"\d+\. ", line):
        paragraph = document.add_paragraph(style="List Number")
        inline(paragraph, re.sub(r"^\d+\. ", "", line))
    else:
        paragraph = document.add_paragraph()
        inline(paragraph, line)
    document.paragraphs[-1].paragraph_format.widow_control = True

document.core_properties.title = "Cosmix Module Integration and Delivery Plan"
document.core_properties.subject = "Review plan for selective module reuse and secure SRS delivery"
document.core_properties.author = "Cosmix Project Planning"
document.core_properties.keywords = "review draft, integration, SRS, inventory, CRM, security"
document.save(OUTPUT)
print(OUTPUT)
print(f"Paragraphs: {len(document.paragraphs)}; tables: {len(document.tables)}")
