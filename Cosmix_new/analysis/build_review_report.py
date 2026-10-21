import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent / 'code'
COMMIT = 'c4500264431523f7db7960bb3a759a15ca60f08e'
OUT = ROOT / 'Cosmix_New_Code_Review_and_SRS_Gap_Assessment.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.7)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(7)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(24)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 3'].font.size = Pt(11)
doc.styles['Heading 3'].font.bold = True

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text, level=1, new_page=False):
    para = doc.add_heading(re.sub(r'\s+', ' ', re.sub(r'[^A-Za-z0-9 ]', ' ', text)).strip(), level)
    if new_page:
        para.paragraph_format.page_break_before = True
    return para

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for cell, label, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = label
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '343434'); cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ix, row in enumerate(rows):
        cells = t.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width); cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ix % 2 == 1:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F4F4F4'); cell._tc.get_or_add_tcPr().append(shade)
    for row in t.rows:
        prevent = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(prevent)
        for cell in row.cells:
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), '85'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            cell._tc.get_or_add_tcPr().append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1.02
                for run in para.runs:
                    run.font.size = Pt(9.5)
    p('')
    return t

def reference(path, line):
    full = REPO / path
    assert full.is_file(), path
    assert 0 < line <= len(full.read_text(encoding='utf8').splitlines()), (path, line)
    para = p('Code evidence  ')
    para.paragraph_format.space_after = Pt(4)
    link = OxmlElement('w:hyperlink')
    rel = para.part.relate_to(f'https://github.com/bkhanzaza551-a11y/cosmix_New/blob/{COMMIT}/{path}#L{line}', RT.HYPERLINK, is_external=True)
    link.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '24516B'); props.append(color)
    run.append(props); text = OxmlElement('w:t'); text.text = f'{"/".join(path.split("/")[-2:])} line {line}'; run.append(text); link.append(run); para._p.append(link)
    for r in para.runs: r.font.size = Pt(9)

def module(name, ids, existing, gap, refs):
    start = len(doc.paragraphs)
    h(name, 2)
    p('SRS coverage  ' + ids)
    p('Present in code  ' + existing)
    p('Remaining workflow  ' + gap)
    for path, line in refs:
        reference(path, line)
    for paragraph in doc.paragraphs[start:-1]:
        paragraph.paragraph_format.keep_with_next = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Cosmix New code review  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
for r in footer.runs: r.font.size = Pt(9)
doc.core_properties.title = 'Cosmix New Repository Review and SRS Gap Assessment'
doc.core_properties.subject = 'Code readiness and reuse decision'
doc.core_properties.author = 'Cosmix Development Review'

p('Cosmix New Repository Review and SRS Gap Assessment', 'Title')
p('Code readiness and reuse decision for the development team', 'Subtitle')
p('12 October 2026')
p('Recommendation  Keep this clone as a candidate foundation, but do not replace the existing system or promise the full SRS in five months on the strength of its current report. The frontend base builds; the backend does not. Many added ERP functions are document storage or status changes, with missing permissions, business handoffs and stock or ledger effects.')
p('The immediate decision is whether this foundation can pass a controlled stabilization and business-flow proof. Adding more repositories before that proof would introduce more integration work without resolving the identified gaps. Three developers are available; the production release scope and team velocity are not yet confirmed.')
p('This assessment combines static source review with actual local build and unit-test checks. Runtime API access, database provisioning, provider integration, browser end-to-end workflows and user acceptance were not verified. Static failure scenarios below are not claims of executed exploits or database incidents.')

h('Review boundaries and source baseline')
p('Reviewed GitHub repository  https://github.com/bkhanzaza551-a11y/cosmix_New.git')
p('Pinned commit  ' + COMMIT)
p('Local code  C:/Users/Huma.Khan/Downloads/Cosmix_Erp_software/Cosmix_new/code')
p('The original 70-section SRS is the historical reference. The current comparison uses the AI-free detailed SRS v0.4 plus HR Attendance and Employee Mobile SRS v1.2. AI assistants, automatic AI classification, generated replies and predictive forecasting are excluded. Required HR face verification for IN is retained. Rule-based routing and calculations are not excluded AI features.')
p('The 17 presentation groups, 23 panels in the supplied report and 96 detailed specification groups use different levels of grouping. Their counts are not interchangeable completion measures. v0.4 includes proposed contract fields and ownership splits; unresolved business decisions are not treated as approved policy or existing API/schema requirements.')
p('The supplied development report was checked as a set of claims. Its migration commands, proposed radii, SLA hours, OTP length and month-five AI work were not executed or adopted. No tracked application/test code or source Word document was changed. Restore, build and test outputs are local generated files.')

h('Verified build and test results', new_page=True)
table(['Check', 'Observed result', 'Meaning'], [
    ('Clone and commit', 'Clone completed and commit pinned', 'A reproducible code snapshot, not a working delivery.'),
    ('Backend solution build', 'FAILED with 16 compile errors', 'Application layer stops compilation; further errors may be masked.'),
    ('Domain unit tests', '2968 passed and 1 failed of 2969', 'Compiled domain suite only. Failed permission-count assertion expects 150 but actual count is 174.'),
    ('Frontend dependency restore', 'Pinned pnpm 10.28.1 frozen restore passed', 'Initial bundled newer pnpm mismatch was an environment issue, not counted as a repository defect.'),
    ('Frontend TypeScript', 'tsc -b passed', 'Static typing of existing frontend source.'),
    ('Frontend production bundle', 'Vite build passed with warnings', 'Circular chunks/reexports and large chunks require runtime review; no ERP business acceptance implied.'),
    ('Frontend unit tests', '140 passed in 11 files', 'Existing unit suite, not the README 154 claim or service/mobile end-to-end coverage.'),
    ('Application and architecture tests', 'Not run successfully', 'Blocked by solution compilation.'),
    ('Relational migrations and integration tests', 'Not executed', 'No database was altered or queried. Current ERP migration gaps are static findings.'),
    ('External providers and UAT', 'Not exercised', 'No Jazz SMS, mail ingestion, WhatsApp, face provider, live GPS or customer acceptance proven.'),
], [1.55, 2.25, 3.3])
p('Commands executed  dotnet build src/NOIR.sln; dotnet test tests/NOIR.Domain.UnitTests/NOIR.Domain.UnitTests.csproj --no-build --no-restore; pinned pnpm install --frozen-lockfile --ignore-scripts; node node_modules/typescript/bin/tsc -b; node node_modules/vite/bin/vite.js build; node node_modules/vitest/vitest.mjs run.')
p('The frontend compiler/bundler/test scripts were invoked directly with the bundled Node runtime after pinned restore to avoid the newer bundled package manager automatically reinstalling dependencies. The source and lockfiles remain unchanged. Lifecycle scripts were suppressed during restore.')

h('First release blockers', new_page=True)
h('Backend compile integration', 2)
p('Fourteen observed CS0104 errors are caused by two PaymentMethod enums imported together. Two CS0535 errors arise because CompleteStockTransferCommand does not implement the audit interface members. These are integration defects in the reviewed commit, not caused by an empty customer database. Fixing them would only unlock the next build; it would not complete the missing business workflows.')
reference('src/NOIR.Application/GlobalUsings.cs', 155)
reference('src/NOIR.Application/Features/Stock/Commands/CompleteStockTransfer/CompleteStockTransferCommand.cs', 6)
h('Committed database provisioning', 2)
p('The App migration snapshot and creation chain omit the new ERP aggregates, including BOQ, journals, stock transfers, production, supplier PO, customer PO, invoices, intake, service and QA records. DbSets and EF configurations exist, but are not an applied or reproducibly provisionable SQL schema. Startup runs migrations and seeding automatically, so the application was not started against the user database.')
reference('src/NOIR.Infrastructure/Persistence/ApplicationDbContext.cs', 160)
reference('src/NOIR.Infrastructure/Persistence/ApplicationDbContextSeeder.cs', 80)
p('The report migration commands also omit --context despite two contexts. The next step is a reviewed migration on an isolated database, after build recovery and schema decisions; not blindly running database update on an existing office database.')
h('Added ERP navigation', 2)
p('Twenty-one advertised added ERP destinations have no corresponding page route. Visible sidebar links fall through to the catch-all redirect. Existing CRM, directory HR and project screens are real; the blanket statement that none of the 23 panels has frontend screens is too broad.')
reference('src/NOIR.Web/frontend/src/components/portal/Sidebar.tsx', 201)
reference('src/NOIR.Web/frontend/src/App.tsx', 233)

h('What is actually in the repository', new_page=True)
inventory = json.loads((ROOT / 'repo_inventory.json').read_text(encoding='utf8'))
table(['Source inventory', 'Count'], [
    ('Tracked files', inventory['trackedFiles']),
    ('Tracked C sharp source files including tests and migrations', inventory['trackedCsFiles']),
    ('Tracked frontend TypeScript and TSX files', inventory['frontendTsTsxFiles']),
    ('Endpoint source files', inventory['endpointSourceFiles']),
    ('Endpoint mapping calls in Program', inventory['endpointMappingCallsInProgram']),
    ('HTTP registration patterns in endpoint source', inventory['httpRegistrationsInEndpointSourceFiles']),
    ('Committed App migration files excluding designers and snapshot', inventory['appMigrationFiles']),
    ('Backend test source files', inventory['testCsFiles']),
], [5.5, 1.6])
p('Counts describe tracked source, not independent APIs, implemented SRS modules, passing cases or a completion percentage. Conditional registrations and multiple routes per file make endpoint and module counts different.')
p('Architecture  .NET 10 and React 19 with TypeScript/Vite; SQL Server/EF Core; Wolverine command/query handlers; repository specifications; Finbuckle tenant filtering; Identity/JWT/cookie/API-key authentication; audit infrastructure; Hangfire jobs; SignalR; media/email/webhook support; a generic PWA shell. These are reusable engineering components subject to the observed failures and release checks.')
p('Existing user-facing base  User/role/tenant/settings administration, contacts/companies/lead pipeline, employee/department/tag/org-chart directory, project/task boards, ecommerce products/orders/customer directory, commerce receipt/history, media, notifications and commerce reports.')
p('Added ERP behavior  BOQ, manual accounts/journals, stock documents and BOM/production have command/query scaffolding with some validations. Procurement, engineering, intake, customer PO, invoicing, service, delivery, helpdesk and QA mainly use direct DbContext endpoints. Direct endpoints are not inherently wrong; the problem is missing equivalent authorization, validation, state and transaction controls.')
p('Definitions  An entity is a stored business-record model. An API is a callable server action. A handler is code that processes an action. A migration creates or changes database tables. A UI route connects a URL to a usable screen. A ledger records authoritative stock or financial effects. Idempotency means retrying the same action does not repeat its effect. These layers must agree before a workflow is accepted.')

h('Department and application gap map', new_page=True)
p('The following 17 groups match the presentation structure. SRS identifiers point to current specification groups, not code module IDs. Present means source behavior was found; no row certifies a production-ready complete module.')
modules = [
('Sales and commercial', 'SAL01 to SAL08', 'CRM contacts, companies, lead pipeline, activities and customer conversion; customer PO header and comparison skeleton.', 'Lead-specific attribution, owned intake handoff, a separately approved and issued quotation/revision, full PO lines and terms comparison, discrepancy resolution and acknowledged Operations handoff.', [('src/NOIR.Web/Endpoints/CustomerPoEndpoints.cs', 53), ('src/NOIR.Application/Features/Crm/Commands/WinLead/WinLeadCommandHandler.cs', 41)]),
('Estimation', 'EST01 EST02 and SAL04 SAL05', 'BOQ header/items, markup arithmetic, submit/approve/reject and total-only revision records.', 'Approved content must be frozen or revised. Full source/equipment/cost versions, alternates, available stock/lead time and approved quotation basis are not completed. Dedicated BOQ workspace is missing.', [('src/NOIR.Application/Features/Estimation/Commands/UpdateBoq/UpdateBoqCommandHandler.cs', 38)]),
('Engineering', 'ENG01 to ENG03', 'Submittal document/revision counter, direct review action and factor-only unit conversion.', 'Technical equipment selection, original/normalized/final values, governed formula versions, AUX definition, submission/review gates and retained approved technical handoff. Celsius/Fahrenheit cannot be represented by factor-only conversion.', [('src/NOIR.Domain/Entities/Engineering/Submittal.cs', 65), ('src/NOIR.Domain/Entities/MasterData/UnitConversion.cs', 39)]),
('Operations and field service', 'OPS01 to OPS08', 'Project/task source and UI; manually assigned FieldJob; site coordinate fields and basic start/complete actions.', 'Accepted-scope project linkage, qualified and available nearest staff selection, offer/accept/decline/timeout/fallback, authorized assigned-work views, live location evidence and operational zones, safe completion and service closure.', [('src/NOIR.Web/Endpoints/FieldJobEndpoints.cs', 81), ('src/NOIR.Domain/Entities/Pm/Project.cs', 8)]),
('Inventory and stores', 'INV01 to INV08', 'Commerce product-variant quantities and receipt/history; stock indent/transfer documents; enterprise GRN model.', 'Authoritative store/location balances, decimal engineering units, accepted receipt effects, issue/return/adjustment approvals, in-transit transfer and destination discrepancy, source-linked aging and valuation. Complete transfer currently changes document status only.', [('src/NOIR.Application/Features/Stock/Commands/CompleteStockTransfer/CompleteStockTransferCommand.cs', 52), ('src/NOIR.Domain/Entities/Product/ProductVariant.cs', 19)]),
('Procurement', 'PR01 to PR08', 'Supplier list/create, PO header/items and approve action; a requisition-like stock indent.', 'Demand/BOQ linkage, approved requirement conversion, supplier/agency governance, quotations/comparative selection and reason, value-based authority, issue/acknowledgement/revisions, partial receipt/closure and procurement reporting.', [('src/NOIR.Web/Endpoints/ProcurementEndpoints.cs', 69), ('src/NOIR.Domain/Entities/Procurement/PurchaseOrder.cs', 74)]),
('Manufacturing', 'MFG01 MFG02 and OPS02', 'BOM validation/scaling and production requirement snapshot, release/start/progress/material-issue/complete handlers.', 'Approved project demand and frozen revisions, material stock debit and returns, lifecycle-safe issue, once-only concurrent progress, applicable QA/rework, finished-goods receipt and controlled handoff. Quantity counters alone do not move stock.', [('src/NOIR.Application/Features/Manufacturing/Commands/IssueMaterial/IssueMaterialCommandHandler.cs', 50), ('src/NOIR.Domain/Entities/Manufacturing/ProductionOrder.cs', 182)]),
('Quality assurance', 'QA01 QA02', 'Inspection record with stage, text checklist, result, defects/actions and evidence URL.', 'Versioned approved criteria, trusted reviewer/evidence, correction and reinspection, linkage to exact source version, applicable receiving/production/project release gates. QA should be applied at agreed checkpoints, not forced onto every transaction.', [('src/NOIR.Domain/Entities/Qa/QaInspection.cs', 67)]),
('Human resources', 'HR01 to HR08 and HR v1.2', 'Meaningful employee/department/tag directory, manager checks, user link, import/export and headcount reports.', 'Expanded private employee fields, controlled activation/offboarding, dated assignments/shifts/holidays, face IN and no-face OUT, raw punches/sessions, overnight/offline/time rules, corrections/month locks, leave/OT, meal/fuel/vehicle/loan effects and locked payroll-input handoff. No attendance engine was found.', [('src/NOIR.Domain/Entities/Hr/Employee.cs', 8), ('src/NOIR.Infrastructure/Persistence/ApplicationDbContext.cs', 135)]),
('Office employee mobile panel', 'EMP01 to EMP06 EMP08', 'Generic signed-in account profile/sessions and shared notifications; responsive admin HR screens.', 'Owned employee profile/document requests, IN/OUT/calendar, correction/site reports, leave/OT, salary effects/claims, and current-team supervisor decisions. A PWA manifest or account profile is not this domain app.', [('src/NOIR.Web/frontend/src/App.tsx', 215), ('src/NOIR.Web/Endpoints/AuthEndpoints.cs', 352)]),
('Field and site employee panel', 'EMP02 EMP03 EMP06 EMP07 EMP08', 'FieldJob server record and manual status API, but no rendered technician destination.', 'Assigned jobs, acceptance/decline, safe job details, duty/site presence and face-IN attendance distinction, evidence/photos/material reporting, sync/retry, trusted recipient proof and appropriate supervisor views.', [('src/NOIR.Web/frontend/src/components/portal/Sidebar.tsx', 232), ('src/NOIR.Web/Endpoints/FieldJobEndpoints.cs', 36)]),
('Customer application', 'CUS01 to CUS07', 'Customer UserId association and administrative directory; complaint list/create/assign/resolve source.', 'Customer-owned complaint/site/evidence submission, receipt/timeline, visit/rescheduling communication, confirm/dispute/authorized no-response closure and own service history. CustomerId from the request is not proof of ownership.', [('src/NOIR.Domain/Entities/Customer/Customer.cs', 15), ('src/NOIR.Web/Endpoints/ComplaintEndpoints.cs', 36)]),
('Accounts', 'ACC01 to ACC08', 'Chart of accounts and manual journal CQRS; invoice/payment models and direct actions.', 'Trusted source/version references, vendor bills/AP, eligible accounts/open periods, exact approved balance/rounding rules, journal approval/reversal treatment, authoritative trial balance and AR/AP aging, closing, HR package verification. Invoice/payment events do not create linked journal entries.', [('src/NOIR.Domain/Entities/Accounting/JournalEntry.cs', 118), ('src/NOIR.Web/Endpoints/InvoicingEndpoints.cs', 73)]),
('Finance', 'FIN01 to FIN08', 'Executive invoice-total KPI and inherited commerce reports.', 'Financial oversight must derive from the single Accounts ledger, with approved margin/tax/currency/rate/value authority and period/HR decisions. Current revenue sum mixes currencies and drops PartiallyPaid invoices. Accounts/Finance ownership requires sign-off; a second ledger is not recommended.', [('src/NOIR.Web/Endpoints/ExecutiveDashboardEndpoints.cs', 21)]),
('Administration', 'ADM01 to ADM08', 'User/role/settings administration and platform tools; active Ticket/comment storage; project tasks.', 'Scoped access-change approval, common reference/approval ownership, trusted helpdesk identity/status/SLA, special task/event workflow and governed imports/reminders. User-supplied ticket/comment author is not trustworthy. A richer unused HelpdeskTicket model must not be counted as API behavior.', [('src/NOIR.Web/Endpoints/TicketEndpoints.cs', 36), ('src/NOIR.Infrastructure/Persistence/ApplicationDbContext.cs', 191)]),
('Superadmin and platform administration', 'SA01 to SA08', 'Platform-admin and tenant-admin roles, permission policy infrastructure, feature definitions, audit/configuration and background jobs.', 'Complete permission registration/enforcement across added APIs, secure bootstrap/secrets, controlled deployment/backup/restore, integration failure ownership and verified recovery. Technical admin is not automatic CEO, payroll or domain approval authority.', [('src/NOIR.Domain/Common/Permissions.cs', 535), ('src/NOIR.Web/Program.cs', 436)]),
('Shared functions and management visibility', 'SH01 to SH10 and common contracts', 'Media, notifications, audit, email templates, outbound webhooks, background-job infrastructure and commerce reports.', 'One governed master/calculation/approval contract, reliable inbound/outbound provider reconciliation, safe files/export scope, immutable business versions, transaction/retry protection, operational location control and traceable ERP/CEO reports. Infrastructure does not supply the missing business engines.', [('src/NOIR.Web/Endpoints/IntakeEndpoints.cs', 26), ('src/NOIR.Web/Program.cs', 814)]),
]
for row in modules: module(*row)

h('Shared controls required across departments')
table(['Shared function', 'Existing foundation and remaining control'], [
    ('SH01 Master data', 'Product/category/brand/attribute basics exist. Store, approved tax/rates, controlled units and effective-date versions are not complete.'),
    ('SH02 Calculations', 'Local quantity-rate/markup/exclusive-tax arithmetic exists. Shared approved tax, WHT, currency, costing and historical inputs are missing.'),
    ('SH03 Approvals', 'Some BOQ/journal/stock/production action permissions exist. No common business authority/value-limit/current-assignment decision contract across added APIs.'),
    ('SH04 Communication', 'Email templates and outbound webhook infrastructure exist. Email/WhatsApp/social/SMS intake connectors, durable source IDs, verified callbacks, unknown-outcome reconciliation and controlled replies are missing.'),
    ('SH05 Files and exports', 'Generic media/file and directory export tools exist. Business-owned evidence/version linkage and HR/customer/private-field access must be enforced for each workflow.'),
    ('SH06 Tasks and notices', 'Project tasks and notifications exist. Assigned service acceptance, special task ownership and required insurance/approval/SLA reminder policies are not complete.'),
    ('SH07 Reports and KPIs', 'Commerce/headcount and project dashboards exist. Authoritative ERP aging/trial balance/payroll/service and drill-down lifecycle metrics are incomplete.'),
    ('SH08 History and audit', 'HTTP/handler/entity logging exists. It does not freeze approved documents, authenticate a request-supplied reviewer, or make retries once-only.'),
    ('SH09 Location and zones', 'Site coordinate fields exist. Fresh staff presence, eligibility/distance routing, reviewed zone events and separate attendance geofence logic are missing.'),
    ('SH10 Platform and recovery', 'Identity/tenancy/configuration/job infrastructure exists. Runtime authorization, fresh SQL provisioning, hardening, deployment, backup restore and release evidence remain unverified.'),
], [1.5, 5.6])
p('Important boundary  Tenant filtering was found and is not dismissed. It isolates modeled tenant data; it does not prove that a signed-in person is the assigned technician, owning customer, current supervisor or authorized financial approver. Dedicated negative role/ownership and tenant-header/JWT mismatch tests remain necessary.')

h('Detailed code findings')
p('P1 identifies a blocker or integrity/access risk for the affected selected workflow. P2 identifies a material correctness or integration gap. Findings below are static observations except for the separately documented build/test results. Repeated concerns across departments are not separate effort estimates. Absence means not found in the reviewed tracked implementation; an external system was not assumed to fill the gap.')
finding_counts = {}
for name, label in [('sales_workflow_ui_review.md', 'Sales engineering projects and UI'), ('finance_supply_review.md', 'Finance supply and manufacturing'), ('people_service_review.md', 'People mobile and service')]:
    h(label, 2)
    source = (ROOT / name).read_text(encoding='utf8')
    matches = list(re.finditer(r'^### ((?:F\d+ \[P[12]\]|P[12]-D\d+:|PS-\d+ \[P[12]\]).*)$', source, re.M))
    finding_counts[name] = len(matches)
    for match in matches:
        end_match = re.search(r'^#{1,3} ', source[match.end():], re.M)
        end = match.end() + end_match.start() if end_match else len(source)
        body = source[match.end():end].strip()
        first = body.split('\n\n')[0]
        summary = first.split('Evidence:')[0].strip()
        summary = re.sub(r'`[^`]+`', lambda m: m.group(0).strip('`') if ':' not in m.group(0) else '', summary)
        summary = re.sub(r'\s+', ' ', summary).strip()
        if not summary:
            continue
        if match.group(1).startswith('P1-D01'):
            summary = 'An authenticated tenant user can reach supplier creation, PO creation and approval, GRN creation, and invoice/payment actions without the corresponding business permission or module gate. PO approval has no trusted reviewer identity, amount limit or decision evidence. Tenant filtering does not supply these within-tenant authority checks.'
        if match.group(1).startswith('F12'):
            summary += ' MoveLeadStage also fails to check that the target stage belongs to the lead pipeline. A direct request can change StageId to another pipeline while preserving the old PipelineId, leaving board and history inconsistent.'
        finding_start = len(doc.paragraphs)
        h(match.group(1), 3)
        p(summary)
        refs = []
        for path, line in re.findall(r'`((?:src|tests)/[^`]+?):(\d+)`', body):
            item = (path, int(line))
            if item not in refs:
                refs.append(item)
        if match.group(1).startswith('F12'):
            refs = [r for r in refs if 'MoveLeadStage' in r[0] or r[0].endswith('/Lead.cs')]
        for path, line in refs[:3]: reference(path, line)
        for paragraph in doc.paragraphs[finding_start:-1]:
            paragraph.paragraph_format.keep_with_next = True

h('Business flow acceptance gaps', new_page=True)
for title, text in [
('Inquiry to accepted execution', 'Retain source/customer/thread -> reviewed intake owner -> validated and versioned BOQ/technical selection -> approved and issued quotation -> complete customer PO comparison -> authorized discrepancy resolution -> acknowledged Operations handoff -> linked execution project. CRM win and matching total do not establish this chain.'),
('Supply to usable stock', 'Approved demand -> approved supplier selection/PO -> expected versus actual receipt -> applicable QA/rejection disposition -> once-only accepted store movement. Receipt header confirmation without the stock effect is not success; multiple item effects must be atomic.'),
('Transfer and production', 'Source availability -> authorized dispatch/in-transit -> actual destination receipt/discrepancy -> conserved balances. Production must tie approved demand and frozen requirements to material issues/returns, progress, applicable QA and accepted finished goods.'),
('Complaint to customer closure', 'Owned customer complaint/site -> eligible available qualified staff before distance ranking -> offer acceptance or next-candidate/supervisor fallback -> assigned work/evidence -> customer confirm/dispute or authorized no-response review -> controlled closure/reopen. A technician Resolve setter is not customer acceptance.'),
('Employee attendance to Accounts input', 'Approved activation/enrollment -> live one-to-one face IN proof -> immutable IN/OUT events with no face for OUT -> sessions/date-effective policy -> separate corrections/leave/OT/approved effects -> reviewed monthly lock -> versioned Accounts handoff. Operational zone warnings remain separate from automatic payroll/discipline.'),
('Financial source to reports', 'Approved billing/receipt/bill/journal treatment -> eligible accounts and open period -> one authoritative linked posting path -> reconciled trial balance/AR/AP views -> authorized closing/correction. Payment approval/settlement and full payroll/bank transfer are only included when management approves that scope.'),
]:
    h(title, 2); p(text)
p('Every accepted route needs authorized and denied API tests, invalid references, legal/illegal transitions, failed-provider outcomes, duplicate/stale requests, concurrent edits and retained evidence. Add actual web/mobile end-to-end tests against a migrated isolated SQL database; inherited unit-test totals cannot substitute.')

h('Corrections to the supplied development report', new_page=True)
table(['Report claim or proposal', 'Assessment'], [
    ('Backend 45 to 50 percent and overall 35 to 40 percent', 'Not validated. No signed requirement weights, acceptance denominator, runtime evidence or current HR/mobile coverage supports these percentages.'),
    ('All 23 panel schemas and APIs completed', 'Models and registrations are partially supported. The build fails, committed ERP provisioning is missing and many APIs only store records or set status.'),
    ('No specialized ERP screens exist', 'Narrow to the 21 added unwired destinations. CRM, base HR directory and PM already have genuine routed screens.'),
    ('34 pages and 25 modals remain', 'A proposal, not an accepted screen inventory. Page count is not workload; customer/employee flows and permission/error states need their own design.'),
    ('Double entry means debit equals credit', 'Current code allows a 0.01 difference and lacks fiscal/account/concurrency controls. An approved rounding and posting policy is required.'),
    ('PO reconciliation compares line items', 'The calling endpoint supplies an empty PO-item list, so same total and currency can produce a misleading Exact Match.'),
    ('Month five AI engine and prediction', 'Excluded by current AI-free scope. Required face IN and deterministic dispatch still remain in scope.'),
    ('200 metre radius and fixed SLA hours', 'Not adopted as company policy. Routing, GPS accuracy, SLA, OTP, authority and no-response decisions need sign-off.'),
    ('Three quotes supplier portal and OCR integrations', 'Conditional proposals rather than automatically approved extra scope. Choose necessary channels and a manual verified fallback before pricing them.'),
    ('Immediately generate and apply migration', 'Unsafe as an instruction for this checkout: repair build, choose the correct context, review schema and use an isolated database first.'),
], [2.5, 4.6])

h('Reuse and delivery decision', new_page=True)
p('Current recommendation  Conditional reuse candidate, not an approved replacement. Retain the NOIR platform/UI patterns and useful CRM/directory/project components as candidates. Treat the new ERP document and service code as unaccepted implementation requiring repair and business validation. Do not merge five or six entire systems or duplicate identities, customer/product masters, stock balances or financial ledgers.')
p('There is no evidence here that borrowing code makes five months impossible, nor evidence that it makes full delivery feasible. The observed gaps make an unconditional full-SRS production promise unsupported. Compile defects may be quick repairs; the missing inventory, finance, attendance and mobile business engines are the larger delivery risk.')
h('Capacity reality for three developers', 2)
p('Using an illustrative 21 working days per month, three full-time developers over five months provide 315 nominal developer-days. The supplied report already totals 277 to 342 developer-days across its 23 estimates. Removing its excluded AI panel leaves 257 to 318 days, without a demonstrated allowance for the broader HR v1.2 gaps, integrations, stabilization, QA, data migration, UAT, deployment or support.')
p('For discussion only, 65 to 75 percent effective feature capacity would be about 205 to 236 days after meetings, review, rework and support. This is a sensitivity assumption, not measured team productivity or a new estimate. Calendar sequencing and three developers sharing specialist domains further limit simple division of effort by three.')
h('First proof before choosing the permanent base', 2)
p('Use a proposed ten-working-day assessment timebox, not a promise to fix every issue in that time. Prove clean backend build and selected suites; reproducible fresh and upgrade SQL migrations; role/owner denial on selected APIs; one atomic supply/stock slice; one owned complaint-to-assigned-work slice; and the architecture fit for HR face/attendance without losing existing work. Record actual effort and remaining blockers. If these proofs fail, re-evaluate the foundation before adding modules.')
p('Compare existing Cosmix HR/backend/mobile work separately before replacing it. Both inspected web project setups are .NET 10, so runtime version alone is not the barrier. Identity/session contracts, tenant versus branch/team scope, employee/customer/product keys, schema ownership, command interfaces and mobile sync are the integration work. Existing repository functionality was not certified by this review.')
h('Rules for selecting any additional donor project', 2)
p('Record source URL, exact commit, original license/NOTICE, files copied, modifications, maintainers and security status. Require a compatible domain model and isolated business logic, not merely similar menu labels. Define one integration contract and migration owner; add positive, negative, retry and concurrency tests. Confirm redistribution and dependency-license obligations before commercial release. A root Apache license does not establish the provenance or licensing of every copied module.')
p('The clone includes an Apache 2.0 license. Apache redistribution conditions include providing the license, retaining applicable notices, marking changes and retaining NOTICE attribution where applicable. Review the original license and donor provenance with the responsible company reviewer. Reference  https://www.apache.org/licenses/LICENSE-2.0')
h('Release hardening and pending evidence', 2)
p('Restore reported package advisory warnings for MailKit, Microsoft.OpenApi, ImageSharp and a test-side SSH.NET dependency. These are dependency metadata findings, not demonstrated reachable exploits. Review runtime reachability and supported patches with regression checks; do not certify security merely because a package comment says an older vulnerability was fixed.')
p('Primary advisory references  https://github.com/SixLabors/ImageSharp/security/advisories/GHSA-j3p4-wp97-rph4 and https://github.com/microsoft/OpenAPI.NET/security/advisories/GHSA-v5pm-xwqc-g5wc')
p('Also resolve public bootstrap credentials/configuration, platform-admin boundaries and production secret injection; prove tenant and same-tenant owner controls. The Docker build stage contains no Node/npm installation while the web publish target invokes npm install/build. That deployment path was not exercised and needs correction/verification. The existing MCP server is an additional integration surface, not evidence of the excluded AI assistant; decide whether to disable it for the agreed release.')
reference('src/NOIR.Web/appsettings.json', 5)
reference('src/NOIR.Web/NOIR.Web.csproj', 29)
reference('Dockerfile', 5)
h('Management decisions that remain open', 2)
p('Confirm mandatory versus deferred modules and native/PWA delivery expectations; company/branch/team authority; store meanings and unit/valuation rules; AUX; approval/value/tax/rate versions; HR conflict resolutions, face provider and retention; geofence accuracy/radius; eligible-staff ranking and workload; complaint SLA/warranty/no-response closure; SMS/WhatsApp/mail accounts and provider terms; fiscal/posting/correction and payroll boundary; data migration and acceptance owners. Do not turn unresolved examples into hardcoded policy.')

h('Evidence record and next review', new_page=True)
p('Original specification  C:/Users/Huma.Khan/Desktop/SOFTWARE REQUIREMENTS SPECIFICATION.docx')
p('Current detailed AI-free specification  C:/Users/Huma.Khan/Desktop/Cosmix_ERP_FSM_Detailed_SRS_v0.4.docx')
p('HR specification  C:/Users/Huma.Khan/Downloads/HR_Attendance_Employee_Mobile_SRS_v1.2.docx')
p('Supplied claim report  C:/Users/Huma.Khan/Downloads/Cosmix_Erp_software/Cosmix_new/HMS_ERP_FSM_Comprehensive_Development_Report.docx')
p('Code anchors in this document link to the reviewed immutable GitHub commit and full source path. Their labels show the owning folder, filename and line. Detailed reviewer notes, extracted source records, source inventory, build logs and Domain TRX result are retained in Cosmix_new/analysis. They support the review but are not runtime/UAT acceptance records.')
p('The next decision should be based on proof results and measured delivery of a selected end-to-end slice. A revised backlog and calendar can then be built for the three developers, with explicit acceptance owners and deferred scope. No application fixes, database migrations or repository merges were performed during this assessment.')

assert len(modules) == 17
assert finding_counts == {'sales_workflow_ui_review.md': 12, 'finance_supply_review.md': 17, 'people_service_review.md': 11}, finding_counts
for root in [doc._element, doc.styles.element]:
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT)
(ROOT / 'report_verification.json').write_text(json.dumps({'commit': COMMIT, 'presentationGroups': len(modules), 'technicalObservationsByReview': finding_counts, 'note': '40 component observations include overlapping controls and are not 40 unique effort items.', 'output': str(OUT)}, indent=2), encoding='utf8')
print(OUT)
