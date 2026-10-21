from pathlib import Path
from copy import deepcopy
import hashlib
import json
import argparse
import re
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
REFERENCE = ROOT / 'deliverables/HR_Backend_Completion_Plan_A_to_G_2026-10-13.docx'
REFERENCE_HASH = 'a90c142b82a9383015dd20a36e7003b5b0ae7460f953b9843cbb376da5023849'
DRAFT = HERE / 'HR_Part_A_Continuation_Status_DRAFT.docx'
FINAL = ROOT / 'deliverables/HR_Part_A_Continuation_Status_2026-10-13.docx'
MARKDOWN = ROOT / 'cosmix-backend/docs/HR-Part-A-Continuation-2026-10-13.md'
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{' + NS['w'] + '}'
parser = argparse.ArgumentParser()
parser.add_argument('--final', action='store_true')
parser.add_argument('--candidate-dir')
args = parser.parse_args()
status = json.loads((HERE / 'status.json').read_text(encoding='utf-8'))
if args.final and not status['final_checkpoint_ready']:
    raise SystemExit('Final checkpoint is not ready; use draft export only.')
if hashlib.sha256(REFERENCE.read_bytes()).hexdigest() != REFERENCE_HASH:
    raise SystemExit('The retained reference changed; inspect before exporting.')

# The complete six-page draft is an intact copied reference package. These
# prefixes identify the eight live-checkpoint paragraphs; no product data is read.
slots = {
    'build': '1. Passed: full solution Release build',
    'independent_review': '2. Five scoped findings',
    'lifecycle_regression': '3. The corrected taxonomy consumer run',
    'model_handoff': '4. PENDING FINAL CHECKPOINT:',
    'private_restore': '1. Passed: 109 existing encrypted objects',
    'database': '1. Passed: fresh current12 COPY_ONLY/CHECKSUM backup',
    'applied_range': '2. Original A23 application passed',
    'http_acceptance': '1. Blocked: the exact LocalTesting.VerifyDatabase',
}
with ZipFile(DRAFT) as zin:
    parts={i.filename:zin.read(i.filename) for i in zin.infolist()}
    infos={i.filename:i for i in zin.infolist()}
tree=etree.fromstring(parts['word/document.xml'])
body=tree.find('w:body',NS)
found=set()
for p in body.findall('w:p',NS):
    current=''.join(t.text or '' for t in p.findall('.//w:t',NS))
    for key,prefix in slots.items():
        if current.startswith(prefix):
            if key in found: raise SystemExit('Ambiguous checkpoint slot: '+key)
            found.add(key)
            text=prefix.split('. ',1)[0]+'. '+status[key]
            runs=p.findall('w:r',NS)
            for r in runs: p.remove(r)
            r=etree.SubElement(p,W+'r')
            t=etree.SubElement(r,W+'t'); t.text=text
if found != set(slots): raise SystemExit('Missing checkpoint slots: '+str(set(slots)-found))

corrections={
    'Overall status: taxonomy/policy source and HR consumers are verified; rehire is source preparation. Local scanner, private restore, SQL restore and selected original A23 upgrade passed. HTTP, genuine policy/provider and rehire activation gates remain; Part A and the whole backend are not operationally complete.': 'Overall status: taxonomy/policy source and HR consumers are verified; rehire is source preparation. Local scanner, private restore, SQL restore and selected original A23/B24 upgrades passed. HTTP, genuine policy/provider and rehire activation gates remain; Part A and the whole backend are not operationally complete.',
    'A-C11 | Status: LOCAL VERIFIED original rollout tracked separately': 'A-C11 | Status: LOCAL VERIFIED SQL restore and selected original migration 24',
    '3. Current live 12/5/14 facts supersede older retained 22/3/11 records. Preserve all original IDs/history/columns; rowversion alone is excluded from content hashes. Backup VERIFYONLY and true isolated restore were both executed.': status.get('current_live_summary', '3. Current live 23/5/14 facts supersede older retained 22/3/11 records; 12/5/14 was the verified pre-upgrade baseline. Preserve all original IDs/history/columns; rowversion alone is excluded from content hashes. Backup VERIFYONLY and true isolated restore were both executed.'),
    'Verify isolated SQL backup restore and frozen A23': 'Verify SQL restore and selected original migrations',
    '4. Finish in order: fresh model/DI and later B24 handoff; guarded acceptance host for the preserved 5-account target; SQL/HTTP retry/concurrency/current-revocation and receipt/audit/decryption/restart checks; genuine policy publication; rehire operational integration; carrier/face acceptance when supplied.': '4. Finish in order: guarded acceptance host for the preserved 5-account target; SQL/HTTP retry/concurrency/current-revocation and receipt/audit/decryption/restart checks; genuine policy publication; rehire operational integration; broader-category Attendance duty/OT integration with Part B; carrier/face acceptance when supplied.',
    '4. Evidence: corrected actual HR consumer run passed 25 checks; old unpinned snapshots remain usable only with exact legacy baseline compatibility.': '4. Evidence: corrected actual HR consumer run passed 25 checks. Old unpinned hiring snapshots remain usable only with exact legacy taxonomy compatibility and absent foundation-field authority.',
    '3. Office requires Branch. Site may be absent in draft/submission but is mandatory at independent approval/materialization. Amendments and Site assignment use dated mapping guards.': '3. Office requires Branch. A Site requirement at activation follows the approved mapping; the original SITE baseline keeps activation Site optional. New mapped categories may explicitly require Site before activation. Scheduled Site duty requires Site. Amendments and Site assignment use dated mapping guards.',
    '5. Next gate: existing-employee category conversion needs a separate approved dated workflow. Attendance roster, office capture and OT consumers need the Part B integration; a category map alone grants no OT, allowance, meal or fuel entitlement.': '5. Next gate: existing-employee category conversion needs a separate approved dated workflow. Broader-category Attendance duty/OT integration remains with Part B; legacy behavior is preserved. A category map alone grants no OT, allowance, meal or fuel entitlement.',
}
for p in body.findall('w:p',NS):
    current=''.join(t.text or '' for t in p.findall('.//w:t',NS))
    if current in corrections:
        for r in p.findall('w:r',NS): p.remove(r)
        r=etree.SubElement(p,W+'r'); etree.SubElement(r,W+'t').text=corrections[current]

if status.get('attendance_source_ready'):
    extra_after=None
    for p in body.findall('w:p',NS):
        current=''.join(t.text or '' for t in p.findall('.//w:t',NS))
        replacement=None
        if current.startswith('Overall status: taxonomy/policy source and HR consumers'):
            replacement='Overall status: taxonomy/policy source and HR plus Attendance consumers are verified; rehire is source preparation. Local scanner, private restore, SQL restore and selected original A23/B24 upgrades passed. HTTP, genuine policy/provider and rehire activation gates remain; Part A and the whole backend are not operationally complete.'
        elif current=='Consume taxonomy in HR':
            replacement='Consume taxonomy in HR and Attendance'
        elif current.startswith('2. Hiring freezes separate joining-date and activation-date version/fingerprint evidence.'):
            replacement='2. Hiring freezes separate joining-date and activation-date version/fingerprint evidence. Review revalidates both; past joining dates may cross a publication boundary with compatible location/schedule semantics.'
        elif current.startswith('3. Office requires Branch.'):
            replacement='3. Office requires Branch. The legacy SITE baseline keeps activation Site optional; new approved mappings may require Site before activation. Scheduled Site duty requires Site. Amendments and assignment use dated guards.'
        elif current.startswith('4. Evidence: corrected actual HR consumer run passed 25 checks.'):
            replacement='4. Old unpinned hiring snapshots remain usable only with exact legacy taxonomy compatibility and absent foundation-field authority.'
        elif current.startswith('5. Next gate: existing-employee category conversion needs'):
            replacement=status['attendance_source_lines'][0]
            extra_after=p
        elif current.startswith('4. Finish in order: guarded acceptance host'):
            replacement='4. Finish in order: guarded acceptance host for the preserved 5-account target; SQL/HTTP retry/concurrency/current-revocation and receipt/audit/decryption/restart checks; genuine taxonomy/policy publication; rehire operational integration; carrier/face acceptance when supplied.'
        elif current.startswith('Evidence locations:'):
            replacement='Evidence: backend docs/Attendance-Taxonomy-Integration-20261013.md, docs/Staff-Taxonomy.md and docs/Hr-Foundation-Policies.md. Scoped reports are in docs/foundation-policies/Verification.md and audit/hr-srs-2026-10-13/part-a-continuation, including roster-taxonomy and taxonomy-overtime verification. The integration report records the consumer counts and points to the 16-file hash manifest.'
        elif current.startswith('2. Evidence: 118 checks passed across 17 explicit consumer gates.'):
            replacement='2. Evidence: 118 pure preparation checks cover 17 named operational gates; their operational adapters and integration remain pending. The real EF model contains none of the new rehire/episode preparation types.'
        if replacement:
            for r in p.findall('w:r',NS): p.remove(r)
            r=etree.SubElement(p,W+'r'); etree.SubElement(r,W+'t').text=replacement
    assert extra_after is not None
    for line in status['attendance_source_lines'][1:]:
        extra=deepcopy(extra_after)
        for r in extra.findall('w:r',NS): extra.remove(r)
        r=etree.SubElement(extra,W+'r'); etree.SubElement(r,W+'t').text=line
        extra_after.addnext(extra)
        extra_after=extra
    for p in body.findall('w:p',NS):
        current=''.join(t.text or '' for t in p.findall('.//w:t',NS))
        if current.startswith('5. Owner: integration lead'):
            extra=deepcopy(p)
            for r in extra.findall('w:r',NS): extra.remove(r)
            r=etree.SubElement(extra,W+'r')
            etree.SubElement(r,W+'t').text='6. Known fixture gap: older VerifyPublication.cs fails to compile with CS7036; its constructor omits IOfficeRosterRepository and OfficeRosterImpactValidator. It is not counted as passed.'
            p.addnext(extra)
            break

# Explicit consumer checklist supplements the short preparation description.
gates='6. The 17 gates are stable identity; previous episode; new approval; current policy; employment terms; required documents; authentication proofs; account reset; fresh access; face isolation; attendance duty; historical coverage; policy/entitlements; financial history; atomic persistence/receipts; legacy reconciliation; employment blacklist.'
for p in body.findall('w:p',NS):
    current=''.join(t.text or '' for t in p.findall('.//w:t',NS))
    if current.startswith('5. Account setup must revoke'):
        extra=deepcopy(p)
        for r in extra.findall('w:r',NS): extra.remove(r)
        r=etree.SubElement(extra,W+'r'); etree.SubElement(r,W+'t').text=gates
        p.addnext(extra)
        break

for p in body.findall('w:p',NS):
    current=''.join(t.text or '' for t in p.findall('.//w:t',NS))
    prefix=re.match(r'^(?:\d\. |A-C\d\d \| Status:)',current)
    if prefix:
        for r in p.findall('w:r',NS): p.remove(r)
        r=etree.SubElement(p,W+'r'); rp=etree.SubElement(r,W+'rPr'); etree.SubElement(rp,W+'b')
        t=etree.SubElement(r,W+'t'); t.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); t.text=prefix.group(0)
        r=etree.SubElement(p,W+'r'); t=etree.SubElement(r,W+'t'); t.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); t.text=current[len(prefix.group(0)):]

parts['word/document.xml']=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
candidate_dir=None
if args.candidate_dir:
    candidate_dir=Path(args.candidate_dir).resolve()
    permitted=Path('D:/CosmixPartAContinuation20261013').resolve()
    assert candidate_dir.is_relative_to(permitted) and candidate_dir!=permitted
    candidate_dir.mkdir(parents=True,exist_ok=True)
    # Preserve the last visually verified pair before constructing a replacement.
    for source,name in [(FINAL,'before-attendance-update.docx'),(MARKDOWN,'before-attendance-update.md')]:
        backup=candidate_dir/name
        if not backup.exists(): backup.write_bytes(source.read_bytes())
    output=candidate_dir/'HR_Part_A_Continuation_Status_CANDIDATE.docx'
else:
    output = FINAL if args.final else HERE/'HR_Part_A_Continuation_Status_REVISED_DRAFT.docx'
output.parent.mkdir(parents=True,exist_ok=True)
with ZipFile(output,'w',ZIP_DEFLATED) as zout:
    for name,data in parts.items(): zout.writestr(infos[name],data)

doc=Document(output)
md=['# HR Part A continuation status','']
for p in doc.paragraphs:
    text=p.text
    if p.style.name=='Title': continue
    if p.style.name=='Subtitle': md += [text,'']; continue
    if p.style.name=='Heading 1': md += ['## '+text,'']; continue
    if p.style.name=='Heading 2': md += ['### '+text,'']; continue
    if re.match(r'A-C\d\d \| Status:',text): md += ['**'+text+'**','']; continue
    if re.match(r'^\d\. ',text):
        md += [text]
    else:
        md += ['',text,'']
mdout=candidate_dir/'HR-Part-A-Continuation-CANDIDATE.md' if candidate_dir else (MARKDOWN if args.final else HERE/'HR-Part-A-Continuation-REVISED-DRAFT.md')
mdout.parent.mkdir(parents=True,exist_ok=True)
mdout.write_text('\n'.join(md)+'\n',encoding='utf-8')

assert len(doc.sections)==1 and len(doc.tables)==0
assert doc.styles['Normal'].font.size.pt==11
assert doc.sections[0].page_width.inches==8.5 and doc.sections[0].page_height.inches==11
assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest()==REFERENCE_HASH
with ZipFile(REFERENCE) as src,ZipFile(output) as dst:
    preserve=[n for n in src.namelist() if n not in ['word/document.xml','docProps/core.xml'] and not (n.startswith('word/footer') and n.endswith('.xml'))]
    assert src.namelist()==dst.namelist()
    assert all(src.read(n)==dst.read(n) for n in preserve)
if args.final:
    assert 'PENDING FINAL CHECKPOINT' not in '\n'.join(p.text for p in doc.paragraphs)
print(json.dumps({'docx':str(output),'markdown':str(mdout),'reference_unchanged':True,'checkpoint_slots':len(found),'body_font_pt':11,'logical_pages':6,'final':args.final},indent=2))
