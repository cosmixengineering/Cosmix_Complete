from pathlib import Path
from datetime import datetime
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'deliverables' / 'HR_Backend_Completion_Plan_A_to_G_2026-10-13.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)

PARTS = []
def part(letter, title, owner, prerequisite, current, anchors):
    p = dict(letter=letter, title=title, owner=owner, prerequisite=prerequisite,
             current=current, anchors=anchors, tasks=[])
    PARTS.append(p)
    return p
def task(p, code, title, kind, srs, work, checks, depends=''):
    p['tasks'].append(dict(code=code, title=title, kind=kind, srs=srs, work=work,
                           checks=checks, depends=depends or p['prerequisite']))

A = part('A', 'Foundation access and employee administration', 'Backend Developer 1',
         'Start first. Agree shared contracts with Developer 2 before parallel consumers.',
         'Authentication, organization masters, hiring, protected employee changes, compensation and limited status transitions already have real services. Complete missing behavior and acceptance rather than rebuilding them.',
         ['src/Cosmix.Application/Authentication/Access/AccessContracts.cs',
          'src/Cosmix.Api/Controllers/Authentication/AccessController.cs',
          'src/Cosmix.Application/HR/Employees/', 'src/Cosmix.Application/HR/Compensation/',
          'src/Cosmix.Infrastructure/Files/', 'src/Cosmix.Api/Controllers/HR/',
          'src/Cosmix.Domain/Common/Approvals/'])
task(A, 'A01', 'Freeze access contracts and complete real authentication acceptance', 'Extension and acceptance',
     'AUTH-001 to AUTH-008; ROLE-001 to ROLE-004; AC-001',
     ['Publish current login/MFA/recovery/device/session/access-assignment DTOs and examples. Retain required Reason in AssignAccessRoleRequest. Assign the initial/additional web payload correction to the web owner; do not relax backend validation.',
      'Add API-003 idempotency receipts for access-creation/assignment writes, which currently use reread reconciliation rather than a replay-key contract. Bind key/payload/actor/action and transaction; reauthorize before replay and publish the changed consumer contract.',
      'Complete approved SMS delivery, secure activation/recovery, rotated refresh tokens, current-state revocation, new-device alerts and first-owner provisioning on the approved target. Reuse existing implementations and never create default production credentials.',
      'Document own, supervisor/team/site, HR, Accounts, fixed Admin and auditor permission/scope boundaries. Recheck live grants and employment status on reads, writes and idempotent replay.'],
     ['Missing/blank Reason is rejected; the corrected initial-access request succeeds once and identical retry returns the same assignment. A different payload using that key conflicts.',
      'Real SMS/MFA, recovery and new-device alerts work; logout, password change and effective suspension revoke access. A supervisor cannot read another site and an employee cannot read another employee.'])
task(A, 'A02', 'Finish private document lifecycle and scanner operations', 'Extension and deployment',
     'HR-002; HR-007; API-005; SEC-001, SEC-007; AC-001',
     ['Retain extension/MIME/signature/size checks, encrypted private storage and scan-before-availability. Deploy approved ClamAV and storage/key configuration; add explicit expire/revoke/archive actions with reason, version and history where absent.',
      'Handle interrupted upload/scan, quarantined or rejected files, retry/rescan, authorization during download and orphan cleanup under approved retention. Keep face storage separate.',
      'Expose required/missing/expiring document and certification metadata to authorized services. Define what blocks activation versus what produces a reminder.'],
     ['A clean verified required document permits hiring submission. Malware, invalid signature, oversized file and revoked user access remain denied, including direct download and retry.',
      'Expired/revoked evidence cannot silently satisfy activation. Upload/download/expire/cleanup decisions have audit records and private objects survive an authorized restore drill.'])
task(A, 'A03', 'Complete employee fields and approved master taxonomy', 'Extension and policy decision',
     'SRS 5.1 and 6.1; HR-002; MD-012',
     ['Add controlled gender/marital values and effective policy validation for age, father/spouse, identity expiry by document type and probation where HR requires them. Save incomplete drafts; enforce approved completeness only at submit/activation.',
      'Resolve the OFFICE/SITE-only category restriction against the SRS driver, rider, purchaser, logistics and field categories. Record approved taxonomy and map scheduling/OT/claim eligibility without duplicating employee records.',
      'Implement permissioned bank verify/revoke transitions with exact verifier, private proof, reason, version and dated history. Finish phone/SIM normalization and activation uniqueness checks; keep bank/medical/salary rights separate.'],
     ['An incomplete draft saves without access. Submission lists exact missing fields; duplicate identity/company email/mobile identity and invalid controlled values fail consistently.',
      'Every approved category resolves documented treatment; category changes retain history. Bank verify/revoke fails for stale or unauthorized actors, records exact evidence and never exposes full bank data through ordinary own/list reads.'])
task(A, 'A04', 'Complete employee lifecycle and archival paths', 'Extension and authority decision',
     'HR-004 to HR-006; AC-009',
     ['Extend the existing lifecycle service for retirement and approved suspension release/reactivation/rehire cases. Define approval authority, effective timestamps and whether rehire appends a new employment episode to the stable employee identity.',
      'Add archive/unarchive and archived-directory filtering where required. Preserve referenced employees, punches, financial evidence, assignments and audit; retain logical archival instead of hard deletion.',
      'Coordinate account/device/session/challenge revocation and biometric retirement at offboarding. Reactivation must not silently restore expired grants or a revoked face enrollment.'],
     ['A future termination blocks punches exactly at the effective instant, including cached tokens and replay, while historical authorized reads still work.',
      'Retirement/rehire/reactivation follow approved transitions, preserve old compensation/status evidence and require explicit current access and face enrollment before Check In.'])
task(A, 'A05', 'Finish post hire assets SIM and certification operations', 'New operational services',
     'SRS 5.1; HR-007; AC-010',
     ['Build issue/return/transfer and condition/custody history for assets and SIMs, with real receiver, date, status and approved credit tier. Reuse hiring intake entities and existing document service.',
      'Add certificate/license renewal, suspension/revocation, replacement evidence and expiry detection. An approved selected-document change must not masquerade as renewal of the certification itself.',
      'Expose current records plus retained history; prevent overlapping custody, double return and silent historical replacement. Feed expiry/custody events to Part G notifications and reports.'],
     ['One asset/SIM cannot have two active custodians. A returned item retains both issue and return evidence; duplicate retry creates no extra custody record.',
      'Renewal preserves the old certificate and attachment; the correct expiry reminder/report changes without exposing medical or identity documents to unauthorized roles.'])
task(A, 'A06', 'Implement required approval routing and delegation', 'New services and scope reconciliation',
     'APR-001 to APR-005; MD-001, MD-005, MD-011',
     ['Implement route evaluation for new claims/benefits using request type, amount/duration, category, department, site and risk where required. Treat existing fixed Admin hiring/leave and HR correction/OT decisions as approved fixed routes.',
      'Add time-bound delegated authority with explicit grant/revoke, scope, permitted action, reason and audit. Delegation cannot grant authority the delegator lacks or allow requester/beneficiary self-approval.',
      'Retain route/authority version, ordered required decisions and before/after evidence. A missing or conflicting route remains pending/unavailable with an actionable reason.'],
     ['Changing a route does not rewrite an existing decision. Threshold/category/site selection chooses one reproducible effective route; ambiguous routes do not auto-approve.',
      'Expired/revoked delegation and self-approval are denied, including retries. Decision notifications contain safe text and retained actor/role/time/reason.'])
task(A, 'A07', 'Coordinate current migrations and shared registrations', 'Integration and acceptance',
     'API-004; SRS 14; AC-010',
     ['Use one migration owner per integration window for DbContext, snapshot, DI and permissions. Review the existing AddOfficeRosterHistory migration and SQL; do not generate a duplicate or rewrite an applied migration.',
      'Compare runtime EF model, snapshot and generated pending SQL. Apply only the reviewed pending range on the separately authorized target, after backup/restore evidence and ownership confirmation.',
      'Validate restrictive relationships, unique active records and schema readiness. Review a data-preserving rollback/forward-fix procedure, including any explicit retention guard required before deployment. Record actual target history and preserve rows/keys.'],
     ['The office migration creates only intended tables/indexes/FKs; the reviewed rollback/forward-fix procedure cannot erase retained history. Model and snapshot agree after coordinated changes.',
      'Approved upgrade preserves existing rows and lets office roster and shared site-shift publication run. The separate receipt/audit/decryption check succeeds without disabling its model guard.'])

B = part('B', 'Attendance evidence scheduling and face verification', 'Backend Developer 2',
         'A01 contracts, A02 storage and A07 schema coordination.',
         'Site scheduling/geofences, office roster source, raw punches/sessions, corrections, face orchestration, Operational/Red zones and site reports exist. Office migration and office leave-source binding now exist in source; full office attendance and live face acceptance remain incomplete.',
         ['src/Cosmix.Application/Attendance/Rosters/', 'src/Cosmix.Application/Attendance/Punches/AttendanceCaptureService.cs',
          'src/Cosmix.Application/Attendance/Calculation/AttendanceCalculationSnapshotProducer.cs',
          'src/Cosmix.Application/Attendance/Corrections/CorrectionReviewService.cs',
          'src/Cosmix.Infrastructure/FaceVerification/', 'src/Cosmix.Application/Attendance/Zones/',
          'src/Cosmix.Application/Attendance/SiteReports/'])
task(B, 'B01', 'Complete office boundary and historical attendance provenance', 'Missing office path',
     'SHIFT-003; ATT-005, ATT-006, ATT-009; OFF-001 to OFF-004; AC-004',
     ['Create approved effective-dated office boundary/location configuration using the real Branch/office identity. Retain exact boundary version and accuracy/radius settings; do not fabricate a site for office staff.',
      'Extend capture, WorkSession revisions, factual/calculated summaries, correction/manual duty and payroll-source snapshots to retain office membership, office roster, shift, calendar and boundary versions.',
      'Use historical assignment at capture/business time and enforce transfer/overnight consistency. Reuse OfficeLeaveScheduleSource and its decision revalidation; do not recreate that implemented binding.'],
     ['An office employee can create an accepted IN/OUT with office proof; missing/inaccurate/outside location has explicit rejected/review status. Site results retain their existing valid source behavior.',
      'An office transfer across overnight duty or a changed retained source produces an explicit conflict/exception. A later boundary edit cannot rewrite accepted evidence.'])
task(B, 'B02', 'Close duplicate sequence and offline intake gaps', 'Extension',
     'ATT-001 to ATT-008; MOB-001 to MOB-005; API-003',
     ['Preserve immutable multi-event punches, authoritative server time and Asia/Karachi business dates. Add approved employee/action/time-tolerance duplicate detection beyond exact retry-key handling without deleting valid break/field events.',
      'Specify offline proof integrity, captured versus received time, idempotency and later validation/review. Queued events remain Unverified until accepted; expired/reused face evidence cannot become accepted merely because an event was offline.',
      'Retain location/device integrity evidence and suspicious-sequence flags. Complete overlap, excessive duration and impossible-travel detection with approved limits and a human review path.'],
     ['Same-key retry returns the same event; changed payload conflicts. A second legitimate break session is retained, while a replay inside the approved duplicate window is detected.',
      'Stale offline face proof, wrong device, mock-location risk and impossible travel never silently produce payable attendance. Received raw evidence and the review outcome remain traceable.'])
task(B, 'B03', 'Complete correction manual duty and geofence exception handling', 'Extension and authority decision',
     'SRS 7.3; ATT-007, ATT-008; FACE-008; AC-003, AC-005',
     ['Extend existing self/supervisor intake and HR review for office manual duty, approved out-of-boundary and accessibility/medical alternative attendance. Define the distinct authority/evidence for each exception; do not treat a general permission as blanket approval.',
      'Retain original punches, requested values, protected evidence, effective sources, independent decision, new session/summary revision and notifications. Make correction result and source trail readable to authorized users.',
      'Route locked-source changes to the first approved later unlocked adjustment period. Preserve the confirmed no-reopen policy and bind the financial effect through Part F.'],
     ['Approve/reject/return/retry preserves raw source bytes and creates only the required revision; a source/version change and requester self-decision fail.',
      'An approved office or alternative-attendance correction has exact source/authority evidence. A locked correction never edits the original lock or produces an unposted monetary effect.'])
task(B, 'B04', 'Deliver the approved live face provider and exception process', 'Provider decision and deployment',
     'FACE-001 to FACE-008; MD-002, MD-004, MD-013; NFR-005; AC-002',
     ['Approve provider/topology, reviewed model/license/hash/input-output semantics, enrollment method, calibrated liveness/match/retry thresholds and exception authority. Current Development-only CompreFace/ONNX PoC is not a production provider.',
      'Implement the approved production adapter and readiness/health/timeout/outage controls while retaining one-to-one enrollment, registered-device/request binding, expiry and atomic single-use proof.',
      'Publish camera-capture/challenge/verification/punch contracts and a documented human alternative route. Face applies only to Check In; Check Out and ordinary app access remain face-free.'],
     ['Genuine live/wrong-person/photo/screen/replay/expired/reused/wrong-device/outage cases are exercised on approved devices. Failure never creates successful Check In or reveals thresholds/templates.',
      'Enrollment/re-enrollment/revocation and verification retain provider/version/actor/device/time/correlation evidence. Check Out succeeds without a face scan.'])
task(B, 'B05', 'Complete biometric retirement deletion and recovery', 'Missing lifecycle operations',
     'FACE-006, FACE-007; SEC-011; MD-013, MD-015',
     ['Add durable post-commit retirement intents for replaced/revoked enrollments, completed captures and offboarding. Wire lifecycle events to retirement; a method existing on a concrete store is not automatic cleanup.',
      'Implement authorized physical deletion and key/backup treatment after the approved retention/hold boundary. Separate expiry of permitted use from authority to destroy stored ciphertext.',
      'Make cleanup idempotent, observable and recoverable after crashes. Validate ACLs, key separation, journal durability and concurrent use-versus-retirement across processes.'],
     ['Revocation immediately denies further use; a crash between transaction and cleanup resumes safely and duplicate cleanup cannot delete an unrelated enrollment.',
      'Approved deletion removes exact live objects and follows approved backup/key policy. No raw capture survives beyond its approved retention without a valid hold.'])
task(B, 'B06', 'Finish zone behavior and reconcile the approved zone scope', 'Extension and SRS deviation decision',
     'ZONE-001 to ZONE-010; MOB-010 to MOB-014; SEC-008A; AC-005A',
     ['Preserve current Operational/Red circle/polygon versions, server assignments, employee precedence, duty windows and separate zone events. Record whether Green/Amber/custom types are formally superseded or still required by v1.2.',
      'If retained, implement Amber permission/exception request/decision and configurable types with their approved authority. Complete safe own instructions, acknowledgements, disputed evidence and reviewer actions.',
      'Define offline event validation/sync and client refresh/window contracts. Monitoring must stop outside authorized duty and use approved recipients, sampling/approach settings and retention.'],
     ['The employee receives only applicable safe instructions; confidential reason fields stay protected. Version/assignment changes do not rewrite old events.',
      'A Red/Amber event remains separate from punches and payroll; unauthorized, out-of-window and stale/offline evidence cannot silently grant access or create a deduction.'])
task(B, 'B07', 'Complete structured site report acceptance and client contracts', 'Acceptance and contract delivery',
     'SRS 7.1; MOB-009, MOB-013; AC-005B',
     ['Reuse existing arrival/departure/delay/access-issue/incident intake, scanned attachments and Accept/Return/Reject review. Document cursor/page shapes, required site/task/time/location evidence and retained response history.',
      'Complete offline/Pending Sync handling, uncertain-network retry, resubmission after return and safe notifications. A site report must not automatically fabricate an accepted punch.',
      'Publish supervisor/site queue discovery and scope-safe detail endpoints where missing. Supply web/mobile owners a runnable contract pack against the actual backend.'],
     ['Online/offline repeated submission creates one report identity and preserves original captured evidence. A different site supervisor cannot view/review it.',
      'Returned/rejected reports retain decisions and resubmission trail. All reporting completes within the application; no external messaging dependency exists.'])
task(B, 'B08', 'Implement or formally close attendance file import scope', 'Missing source workflow and scope decision',
     'SRS 7.1 File import; ATT-001, ATT-003; API-003',
     ['Confirm whether authorized attendance file import remains a release requirement. If superseded, retain an approved scope record; otherwise build validated batch upload/mapping/preview/commit with file hash and importer.',
      'Retain row errors, source evidence, duplicate report and audit. Imported rows cannot impersonate a live mobile face-verified Check In or erase an existing raw event.',
      'Make commit retry-safe and permission-scoped; use approved review/correction for any attendance interpretation. Do not add attendance-machine integration.'],
     ['The same file/batch/key cannot create duplicate records; invalid rows have explicit errors and a denied employee/site cannot be imported.',
      'Committed evidence retains batch/hash/row/actor; original punches are unchanged and no file content bypasses approved mobile/exception verification.'])

C = part('C', 'Effective policies calculations and overtime', 'Developer 1 owns configuration and Developer 2 owns calculations',
         'A03 categories, B01/B02 factual sources and agreed versioned policy contracts.',
         'Typed ordinary time, late/half-day/hourly calculations and confirmed site OT formula exist. KPI, waiting, studying/make-up, approved early exceptions and special continuous-duty policy producers remain incomplete.',
         ['src/Cosmix.Application/HR/Policies/', 'src/Cosmix.Domain/HR/Policies/',
          'src/Cosmix.Application/Attendance/Calculation/', 'src/Cosmix.Application/Attendance/Overtime/',
          'src/Cosmix.Domain/Attendance/PayrollInputs/AttendancePayrollInputRules.cs'])
task(C, 'C01', 'Extend typed effective policy configuration', 'Extension',
     'HR-004; SHIFT-003, SHIFT-004; API-004; AC-004',
     ['Add explicit typed schemas for the new rule families instead of arbitrary JSON or UI constants. Each schema declares applicability, effective period, precision, authority/evidence and validation errors.',
      'Preserve direct HR publication and retained immutable versions. Reuse existing policy resolution and reject conflicting or unapproved configuration with unavailable reasons.',
      'Define version/fingerprint inputs shared by leave, claims, KPI and finance. Locked calculations retain their exact source policy and cannot be silently recomputed under a later version.'],
     ['Same dated sources/configuration reproduce the same result/fingerprint. Unknown fields, conflicting scope and missing approval authority fail explicitly.',
      'A new future rule changes only its intended dates; a locked old period remains byte/source compatible.'])
task(C, 'C02', 'Finish office and site ordinary time treatment', 'Extension and acceptance',
     'OFF-001 to OFF-004, OFF-007, OFF-009; SITE-001 to SITE-004; ATT-009',
     ['Apply current effective office/site schedule and exact approved leave/travel/holiday evidence. Retain scheduled, worked, payable, late, shortage, exceptions and calculation version with a clear explanation.',
      'Preserve approved BASIC/actual-calendar-days ordinary monetary inputs and larger-only overlapping ordinary/late deduction. Bind office sources from B01 instead of accepting missing site provenance.',
      'Validate 10:00 start, 10:15 late threshold, 12:59 half-day/hourly boundary and configured nine-hour/six-day/break rules against signed policy and approved examples. Clarify ambiguous arrival bands before activation.'],
     ['Approved examples at threshold equality and one tick/minute after it calculate correctly; missing OUT/open sessions are not silently payable.',
      'A temporary approved 11:00 site start, office transfer, paid leave/holiday and source correction use the right dated schedule without changing a locked result.'])
task(C, 'C03', 'Add approved early reporting and studying assignments', 'Missing policy producer and decisions',
     'OFF-002, OFF-010, OFF-011; SITE-002; MD-012',
     ['Maintain named/role early-credit exemptions as effective-dated employee policy assignments with written instruction, authority and eligible interval. Default pre-start time remains excluded.',
      'Implement dedicated studying/university schedule, hourly and make-up treatment, TDA exclusion and KPI interaction only from approved rules and examples.',
      'Record which actual work intervals satisfy make-up obligations and prevent double use for ordinary pay, overtime, leave recovery or other credits.'],
     ['An uninstructed early IN earns no pre-start credit; the approved exact interval earns only its permitted treatment.',
      'Study/make-up hours use the retained assignment and cannot satisfy two obligations or silently create TDA/OT/KPI money.'])
task(C, 'C04', 'Implement finalized KPI and annual late assessments', 'Missing engine and decisions',
     'OFF-006; SITE-010, SITE-011; LEAVE-008; MD-009',
     ['Define approved KPI base, year/reset, qualifying events, tier interpretation, precision, bounds and composition. Retain SRS 30/31-60/61-120 boundaries and 0.8/0.5 percent values without guessing their financial meaning.',
      'Build dated factual intake, review/dispute, exact source calculation and authorized finalized-score revisions. Distinguish monthly committed-hours shortage, missing-evidence and late effects.',
      'Expose finalized KPI evidence for D02/D03. A 120-late threshold creates a reviewed disciplinary/escalation case; it cannot automatically terminate or blacklist the employee.'],
     ['Approved cases at counts30,31,60,61,120,121 and across annual reset retain exact composition/rounding and reviewer evidence.',
      'An amendment preserves the original score and cannot create salary deductions or benefits from an unfinalized/display-rounded score.'])
task(C, 'C05', 'Implement paired worker waiting attribution', 'Missing engine and decisions',
     'OFF-005; SITE-005; MD-010',
     ['Create effective work-pair/task dependency records and factual waiting intake with timestamps, evidence and disputes. Do not fabricate a verified actor/approval merely to satisfy an existing model constructor.',
      'Implement independent verification and approved causation/attribution/cap/overlap rules. Distinguish delaying worker time, waiting worker time and productive alternate work.',
      'Only a separately approved financial/performance interpretation produces an effect, retaining pair/source/policy/decision. Pass a proposed financial source to Part F; factual waiting alone is not a payroll charge.'],
     ['A waits10:00-10:20, then works while B arrives10:45: only the approved proven interval is attributed. Overlapping waits from A/C do not multiply one loss silently.',
      'Approved changed start time, excused lateness, disputes and later review change only a retained revision/forward adjustment; self-verification and unapproved pay impact fail.'])
task(C, 'C06', 'Finish continuous duty travel and disciplinary cases', 'Missing special policies and decisions',
     'SITE-006, SITE-009, SITE-010, SITE-011; OT-003; MD-006',
     ['Implement the approved 10-hour reset, next-day4.5-hour and12-extra-hour examples using continuous actual sessions, source dates, breaks/home visits and retained roster context.',
      'Create company-work travel treatment as paid normal work with approved deployment/trip evidence; personal hometown leave travel remains leave/absence. Travel does not automatically earn OT.',
      'Add evidence/review/decision lifecycle for fake IN/OUT and annual-late escalation. Store outcomes and audit without automatic termination, blacklisting or financial punishment.'],
     ['Overnight/duty-reset examples and one-minute boundary cases match approved expected outputs without double counting across business dates.',
      'Company versus personal travel and alleged fraudulent evidence have explicit approved treatment; unresolved evidence cannot become automatic payable time or punishment.'])
task(C, 'C07', 'Complete OT eligibility calculation and financial source delivery', 'Extension and acceptance',
     'OT-001 to OT-005; SITE-008; MD-007',
     ['Reuse named bulk review, exact approved intervals and the confirmed BASIC/source-month-calendar-days/nine-hours formula with final whole-rupee rounding. Office OT remains ineligible.',
      'Resolve approved categories/day types and supported schedules explicitly. Current nine-hour-only behavior must either remain an approved limitation or support additional approved schedule/rate contracts.',
      'Produce immutable approved OT award/evidence for F02, including exact ticks, compensation/policy revision, job/site and authorization. Calculation approval is separate from posting/settlement.'],
     ['October31-day and February28/29-day cases use the source month; fractional seconds survive to final rounding. Split sessions exclude non-work gaps.',
      'Named bulk approval cannot silently affect all staff; overlap/travel/ineligible office cases fail. Finance consumes each retained award at most once without repricing it.'])

D = part('D', 'Leave benefits encashment and recovery', 'Backend Developer 1',
         'A02 evidence, B01 scheduling, C01 policies and C04 finalized KPI for benefits.',
         'Ordinary leave types, July-June balances, evidence, direct Admin decisions, cancellation and exact allocation already exist. New office leave-source retention exists. Complete contract/acceptance gaps and the missing benefit/recovery engines.',
         ['src/Cosmix.Application/HR/Leaves/', 'src/Cosmix.Domain/HR/Leaves/',
          'src/Cosmix.Application/HR/Leaves/OfficeLeaveScheduleSource.cs',
          'src/Cosmix.Infrastructure/Repositories/HR/', 'src/Cosmix.Api/Controllers/HR/'])
task(D, 'D01', 'Finish ordinary leave intake calendars and acceptance', 'Extension and acceptance',
     'LEAVE-001 to LEAVE-010; SRS 9.1; ROLE-002, ROLE-003',
     ['Complete date-range preview/duration/holiday explanation and policy-required contact, handover/reliever fields. Provide own history/balance and authorized team calendar; retain exact historical day/minute denominators.',
      'Preserve fiscal July-June18 days, sick6/casual12, minimumfour-month eligibility where applicable, casual advance approval, medical/exam evidence, injury non-consuming rest and Admin review/cancellation.',
      'Exercise office source freeze/revalidation already implemented and the leave allocation consumer. Define the missing-evidence emergency half-day review/KPI branch through C04 rather than silently approving without evidence.'],
     ['Fiscal rollover, exactlyfour-month service, partial-day, paid holiday, overlapping leave, cancellation and concurrent balance usage have reproducible outcomes.',
      'Office roster/membership changes conflict at review; injury rest does not consume regular leave. Own/team data remain scoped and missing evidence follows an explicit permitted decision.'])
task(D, 'D02', 'Implement extra six days or cash election', 'Missing service and decisions',
     'LEAVE-011; MD-011; MD-009',
     ['Create discretionary management grant and election using an exact finalized100% KPI revision, eligibility period, authority and approved expiry/valuation rules.',
      'Keep the benefit pool separate from ordinary18-day/sick/casual entitlement. Add real balance/query/reservation/usage semantics; inserting an unused Entitlement ledger row is insufficient.',
      'Prevent receiving both full days and full cash for one grant. Bind cash election to an approved F financial source; retain amendment/reversal and already-used benefit history.'],
     ['99.995 displayed as100 does not qualify unless the approved precision rule says so; authoritative100 qualifies only with required grant.',
      'A six-day grant/election retry is unique; switching election after partial usage follows approved rules and cannot duplicate benefit or erase posted history.'])
task(D, 'D03', 'Implement leave encashment with reservation and settlement', 'Missing service and decisions',
     'LEAVE-012; MD-011; AC-007',
     ['Calculate eligible unused regular leave under the approved cutoff/balance rules and finalized KPI of at least70%. Define valuation base/date/divisor/fraction/currency/rounding separately from absence and OT formulas.',
      'Build request/management decision/reserve/release/consume lifecycles with source balance/policy/KPI/compensation fingerprints. Prevent simultaneous leave or another encashment consuming the same units.',
      'Generate approved financial lines for F; support rejection/withdrawal, partial posting and reversal without declaring the entire claim settled from a handoff alone.'],
     ['KPI69.99/70.00/70.01 and1.5-day cases use approved exact precision/day units. Pending leave and the extra-six-day pool are treated explicitly.',
      'Concurrent leave/encashment attempts cannot overconsume balance; partial Finance receipt settles only its exact lines and retains unconsumed/reserved remainder.'])
task(D, 'D04', 'Implement long leave work recovery commitments', 'Missing service and decisions',
     'LEAVE-013; OFF-011 where linked',
     ['Define approved pre-leave/post-return obligations, minute units, deadlines, eligible work, authority, cap, waiver and financial interpretation.',
      'Build commitment/request/approval/fulfillment allocation/dispute/waiver/close operations retaining actual dated work sources. Never use a generic model row as proof that hours were recovered.',
      'Prevent the same interval fulfilling make-up, recovery and another payable credit twice. Route only an approved monetary interpretation through F and keep locked source periods unchanged.'],
     ['A commitment fulfilled across dates retains exact credited intervals and remaining minutes; deadline/waiver/transfer cases follow approved rules.',
      'Duplicate retry or overlapping fulfillment cannot overcredit; missing work proof or an unapproved financial rule remains unavailable instead of becoming a deduction.'])
task(D, 'D05', 'Complete leave audit notifications and downstream contracts', 'Integration',
     'APR-003, APR-005; REP-002; SRS 9.1 and12; AC-010',
     ['Emit safe request/decision/return/reject/cancellation/benefit/encashment/recovery events through G notifications. Expose scoped status/evidence/history without medical payloads in messages.',
      'Publish exact ordinary/benefit ledger and allocation contracts to calculations, ESS, reports and finance; document which transition reserves, consumes, releases and posts each source.',
      'Expose explicit incomplete/review_required state for missing policy/KPI/financial source; a zero/empty result must not imply entitlement or settled liability.'],
     ['The same committed decision emits one intended notification and each consumer uses its retained revision.',
      'Ordinary cancellation and financial benefit reversal retain different controlled actions; report/ESS totals agree with authoritative ledger and exact receipts.'])

E = part('E', 'Allowances vehicles advances and loans', 'Backend Developer 1 with Developer 2 source-contract review',
         'A06 approvals, B factual sessions/reports, C01 policies; agree F01 financial source contract first.',
         'Existing allowance, trip, fuel, service, custody and loan entities are persistence foundations. Operational eligibility, request/review/settlement APIs and financial self-service are missing.',
         ['src/Cosmix.Domain/HR/Allowances/', 'src/Cosmix.Application/HR/Allowances/',
          'src/Cosmix.Domain/HR/VehicleFuel/', 'src/Cosmix.Application/HR/VehicleFuel/',
          'src/Cosmix.Domain/HR/AdvancesLoans/', 'src/Cosmix.Application/HR/AdvancesLoans/'])
task(E, 'E01', 'Build meal and outstation eligibility cycles', 'Missing engine and decisions',
     'ALLOW-001 to ALLOW-004; OFF-008; SRS10.1; MD-008',
     ['Implement continuous-duty eligibility from verified sessions and approved deployment/category: office13-hour duty from10:00 up toPKR400; stationed dinner after12 continuous hours; breakfastPKR300 eight hours later; lunchPKR500 eight hours after breakfast eligibility.',
      'Implement approved outstationPKR1000/day with no additional dinner duplication. Keep unspecified dinner amount and cycle/home-break definitions unconfigured until approved.',
      'Retain cycle/session/deployment/policy/eligibility evidence and verified break closure. Separate eligibility, requested amount, approved amount, payable date and Finance settlement.'],
     ['Boundary minute, home visit, split duty, overnight and outstation cases choose the right cycle; one cycle cannot claim duplicate meal entitlements.',
      'Unknown dinner amount/deployment/continuity remains explicitly unavailable. Approved eligibility does not claim payment without an exact receipt.'])
task(E, 'E02', 'Deliver allowance claims decisions and settlement sources', 'Missing operational services',
     'ALLOW-001 to ALLOW-004; APR-001 to APR-003; API-003',
     ['Create scoped eligibility/claim/draft/submit/review/return/reject/cancel APIs with scanned evidence, amount constraints, versions, retry keys and approved route.',
      'Retain before/after proposal and independent decision, payable period and original duty cycle. Returned resubmission preserves original evidence and resolves duplicate entitlement.',
      'Deliver approved financial source to F01 and own status/history/report inputs to G. Amendments after posting use a new controlled adjustment, not overwriting the claim.'],
     ['Self-approval, duplicate cycle, wrong employee/site and stale version fail. Same-key success retry returns the same claim/decision.',
      'Rejected/returned/unposted claims do not enter settled totals; each approved cycle source can be consumed exactly once.'])
task(E, 'E03', 'Build route mileage commute and fuel reimbursement', 'Missing engine and services',
     'VEH-001 to VEH-003, VEH-006, VEH-007; SEC-009',
     ['Implement approved complaint/field eligibility, effective commute baseline, authorized trip/route evidence and approved site-to-site distance. Exclude home-office-home distance and apply confirmed extra/nearer-distance monthly rules.',
      'Calculate approved fuel quantity using28km/liter; freeze route/map distance, claim period, policy, actual proof and approved financial value. Configure provider/privacy and monetary valuation rather than guessing fuel price.',
      'Reject excluded puncture/fines/lifting/bearing/helmet/seat/cover/indicator expense categories. Expose adjustments, reviewer decisions and immutable approved-mileage ledger.'],
     ['28 approvedkm yields1 eligible liter; an ordinary commute yields no reimbursable distance. Route/commute replacement cannot reprice an approved old claim.',
      'Wrong category, repeated trip, excluded expense and identity-bearing map requests are rejected; monthly excess/nearer adjustment matches approved examples.'])
task(E, 'E04', 'Implement motorcycle maintenance and workshop eligibility', 'Missing engine and services',
     'VEH-004, VEH-005, VEH-009',
     ['Accumulate only approved outstation kilometers in an immutable consumption-aware mileage ledger. Trigger tuning eligibility each1000km atPKR1800 under the effective policy.',
      'Implement overhaul eligibility after8000 approved outstationkm with the approximatelyPKR18000 signed-policy reference and explicit approved claim/cap treatment; do not turn an approximate amount into automatic payment.',
      'Build service request/review/workshop/odometer/evidence/history and consumed-milestone records. Coordinate claim financial sources and vehicle custody.'],
     ['999/1000/1999/2000km tuning and7999/8000km overhaul cases match approved eligibility; duplicates and mileage reversals cannot pay one milestone twice.',
      'Unauthorized workshop, inconsistent odometer and unapproved distance remain review/rejected. Service/payment status is supported by actual source and receipt.'])
task(E, 'E05', 'Complete vehicle custody incidents insurance and liability', 'Missing operational services and authority',
     'VEH-008, VEH-009; SRS5.1',
     ['Implement issue/return/transfer with condition, receiver, effective time and custody history. Connect insurance/license/expiry documents through private storage.',
      'Build accident/third-party/evidence/insurance-claim and reviewed liability workflows with separately approved permissions and policy. An incident record alone cannot impose salary recovery.',
      'Publish custody/incident/service history and reminder/report events. Send approved liability/credit source revisions to F; posted/settled status comes only from the exact authorized Finance receipts.'],
     ['Double custody/return conflicts; an issue followed by transfer/return preserves all condition evidence.',
      'An incident with missing evidence or unapproved liability cannot create payroll recovery. Restricted third-party/medical/identity fields stay scoped.'])
task(E, 'E06', 'Build advances loans repayment and amendments', 'Missing operational services and terms',
     'SRS1.2,11,12,14; AC-006, AC-007',
     ['Approve eligibility, principal/currency, installments, repayment timing, outstanding balance, disbursement evidence and amendment/partial/reversal authority. Do not invent interest or repayment rules.',
      'Build own request/history, scoped review, agreement, disbursement receipt, scheduled installments and actual repayment ledger. Separate approved agreement from money disbursed.',
      'Produce only due approved recoveries for F and own outstanding/history for G. Handle suspension/termination, early repayment, rejected/cancelled request and amendment without erasing prior postings.'],
     ['Duplicate request/receipt cannot double principal or recovery; partial repayment updates the correct installment and outstanding amount exactly once.',
      'An undisbursed agreement never creates recovery. Concurrent amendment and repayment use versions; employee only sees their own financial records.'])

F = part('F', 'Financial consumption payroll handoff and month closure',
         'Developer 1 owns financial sources and Developer 2 owns attendance locks',
         'Agree F01 early; finalize after approved C/D/E sources and B historical evidence.',
         'Ordinary/late saved inputs, exact protected revisions, no-charge handoff and manual Finance receipts exist. Approved OT consumption, monetary carry-forward, new claims/credits, formal legal activation and successful full month closure remain incomplete.',
         ['src/Cosmix.Application/Attendance/PayrollInputs/',
          'src/Cosmix.Application/Attendance/PayrollInputs/AttendanceMonthLockFinancialReader.cs',
          'src/Cosmix.Application/Attendance/MonthLocks/', 'src/Cosmix.Application/Attendance/CarryForward/',
          'src/Cosmix.Application/Attendance/LegalReview/', 'src/Cosmix.Application/HR/PayrollHandoff/'])
task(F, 'F01', 'Freeze shared financial source and consumption contracts', 'New shared integration contract',
     'AC-007; SRS14,15; MD-014',
     ['Define one source-revision contract consumed by handoff: employee, source type/ID/revision, source period, approved policy/decision, exact amount/currency/direction, fingerprint and current availability. Adapt existing ordinary/late/OT precision; do not reprice.',
      'Define durable source-line consumption and release/reversal identifiers. A source consumed in one batch cannot appear again in another batch through retry, report export or a new source alias.',
      'Document request/response states and schema versions shared by C/D/E/G. Every producer distinguishes unavailable, proposed, approved and posted/settled evidence; map existing states explicitly.'],
     ['Two batches racing for one source can create only one valid consumption. Different revisions/currency/amount/decision produce an explicit conflict.',
      'An unapproved or unavailable source cannot become a zero-value liability or a handoff-ready line. Retained old schema sources remain readable.'])
task(F, 'F02', 'Consume approved OT claims and benefit financial sources', 'Missing integration',
     'OT-002, OT-005; ALLOW-003; LEAVE-011, LEAVE-012; AC-007',
     ['Connect exact approved OT awards and approved meal/fuel/service/benefit/encashment/loan-recovery sources to the financial producer with source trace IDs and consumption rules.',
      'Preserve confirmed ordinary/late larger-only overlap behavior. Approve and implement overlap/composition with waiting, recovery, benefit and credits separately; do not borrow a salary divisor without authority.',
      'Provide explicit pending/blocked reasons when any required producer or policy is missing. Keep salary and bank changes separate from posting responsibility.'],
     ['OT approved in a31-day source month uses its retained award even if handoff occurs in another month or salary changes later.',
      'Approved claims appear once with their exact direction/amount; rejected/unposted/reversed sources follow their documented state and cannot be silently settled.'])
task(F, 'F03', 'Complete financial forward adjustments and reconciliation', 'Missing monetary extension',
     'ATT-010; SRS7.3; AC-007',
     ['Convert authorized factual locked-period corrections into approved monetary delta sources with original frozen evidence, current decision and first later unlocked target.',
      'Extend receipt-backed reconciliation for actual credits/recoveries/partial reversals, preserving prior posted/settled evidence. Saved/handoff-only inputs are not posting proof.',
      'Prevent double application across months and keep revision/amendment/reversal trails. Source lock manifests remain immutable; no reopening or old-row deletion is introduced.'],
     ['A late correction to a locked month produces one approved forward delta in the first eligible later month; duplicate retry or a changed target cannot apply it twice.',
      'Missing prior posting, partial receipt and reversed receipt expose explicit incomplete state and reconcile only their exact confirmed amounts.'])
task(F, 'F04', 'Implement formal applicability and legal activation evidence', 'Missing authority contract and operational acceptance',
     'MD-014, MD-015; AC-007; current legal activation gate',
     ['Obtain formally approved authority, applicable worker/location/rule scope and required review evidence for payroll activation. Reuse existing recorded legal dossiers/jurisdiction versions; uploaded PDF or technical permission alone is not legal clearance.',
      'Implement versioned activation/revocation/expiry/scope decision records and server validation tied to exact financial source periods/rules. Do not remove the current unconditional gate without a supported authoritative replacement.',
      'Keep observable unavailable/review-required reasons and block financial release when approval is missing, revoked, expired or contradictory. Management/legal owner supplies applicability decisions, not developers.'],
     ['Valid retained authority/applicability permits the approved financial readiness path; record-only or wrong-location/rule/time evidence remains blocked.',
      'Revoking/changing applicability before finalize or replay invalidates readiness and cannot be bypassed using an old client fingerprint.'])
task(F, 'F05', 'Complete successful month finalization without a dependency cycle', 'Missing positive closure path',
     'ATT-010; SHIFT-004; AC-007',
     ['Define the approved order of final source calculation, exact input revisions, pre-lock handoff/consumption reservation, applicability proof and atomic freeze. Accounts posting follows its authorized receipt path; do not require a lock that cannot yet exist to prepare its own inputs.',
      'Replace current unavailable dependencies only with verified source/authority readers. Finalize one complete nonempty population with exact membership, summaries, adjustments, inputs and fingerprints.',
      'Retain no-reopen behavior, optimistic concurrency and replay reauthorization. Missing/changed evidence must explain why the period is not ready.'],
     ['A complete approved nonempty office/site month reaches one final lock; a zero-charge employee uses the existing no-charge handoff successfully.',
      'One missing OUT, legal source, claim interpretation, pending correction or concurrent source change prevents finalization. Lock retry cannot freeze different evidence or reopen history.'])
task(F, 'F06', 'Complete Accounts receipts partial posting and reversal', 'Extension and actual acceptance',
     'AC-007; MD-014; SRS15 Payroll handoff',
     ['Expose authorized locked-input/batch reads and exact line-level handoff/posted/settled receipt operations. Match revision, line IDs, amount, currency, external evidence and actor authority.',
      'Support documented partial posting, rejection, settlement, amendment and reversal with immutable receipt-linked reconciliation; do not mark a whole batch posted from one partial receipt.',
      'Accounts cannot modify raw punches, policies, original compensation or locks. First release supports approved handoff/bank-advice output where required; automatic bank disbursement is out of scope.'],
     ['An exact authorized receipt changes only its lines. Wrong amount/revision/currency/duplicate or forged source receipt fails and is audited.',
      'Partial/reversed settlement totals reconcile to approved source consumption and own financial views. Accounts write attempts against punches/policies return denial.'])

G = part('G', 'Self service reporting notifications and release acceptance',
         'Developer 1 owns HR consumers and Developer 2 owns jobs and attendance consumers',
         'Start reusable infrastructure/contract work early; finish financial reports after C/D/E/F producers.',
         'Some own attendance/leave/OT/inbox APIs and genuine factual CSV/XLSX/PDF export infrastructure exist. Full own profile/financial services, HR notification producers, all fourteen reports and complete deployed acceptance remain incomplete.',
         ['src/Cosmix.Application/HR/EmployeeSelfService/', 'src/Cosmix.Api/Controllers/EmployeeSelfService/',
          'src/Cosmix.Application/HR/Reports/', 'src/Cosmix.Application/Attendance/Reports/',
          'src/Cosmix.Application/Common/Notifications/', 'src/Cosmix.Api/BackgroundJobs/',
          'src/Cosmix.Api/Authentication/AttendanceTrafficControls.cs'])
task(G, 'G01', 'Deliver own profile home and financial self service APIs', 'Missing API producers',
     'ROLE-003; MOB-007, MOB-008; AC-006; SRS11',
     ['Add authenticated own approved profile/emergency/document summary and permitted edit-request APIs. Use employee identity from the session, not a trusted client employee ID, and mask restricted fields.',
      'Deliver own salary/pay summary, deductions, approved allowances, advances/loans/outstanding/repayment, payslip references and pending/posted/settled states from authoritative producers.',
      'Provide home/next-action/current shift/site/today/pending-request contracts by composing existing services. Never imply finance readiness or fabricate values because a producer is unavailable.'],
     ['Office and site employees see only their own data, including direct URL/query tampering. Medical/department-wide administration is absent from employee projections.',
      'Own totals match exact retained approvals/receipts; unavailable payslip or unposted claim is explicitly labeled and does not look paid.'])
task(G, 'G02', 'Complete scoped supervisor discovery and operational views', 'Missing consumer queries',
     'ROLE-002; API-002; SRS2 and9.1',
     ['Provide current/historical authorized team/site employee discovery for attendance queues, reports and leave calendars. A caller supplies filters, not unrestricted employee IDs as authorization.',
      'Reuse server data scopes at collection/detail/evidence/export time; define query pagination/cursor, sorting and filters consistently for all web/mobile consumers.',
      'Keep material/tool request access a separate permission; supervisor status does not grant department HR or salary/medical rights.'],
     ['A supervisor transfer changes current discovery while retained source evidence uses its required historical authorization; unauthorized IDs never leak in totals or counts.',
      'Pagination/filter changes do not widen scope; direct sensitive-data/detail/export requests fail without separate rights.'])
task(G, 'G03', 'Deliver all fourteen source backed reports', 'Missing report producers',
     'REP-001 to REP-003; SRS15; AC-008',
     ['Extend current real private CSV/XLSX/PDF job generators for the fourteen report catalogue entries below. Register each report against a real producer and expose unavailable/missing source reasons.',
      'Include generation time, exact filters, policy period/version, source/lock state and authorized columns. Sensitive employee/financial/location columns are permission-specific and excluded by default.',
      'Reauthorize requester/session and original/current source scopes at generation/status/download. Implement configured large-report background execution, bounded queries, expiry and audit.'],
     ['Each report creates an actual readable file with totals traceable to its source. User without salary/GPS/medical permission gets neither hidden columns nor leaked metadata.',
      'Revoked grants/session before generation/download deny access; missing financial/KPI producer is not exported as zero or a UI-only success.'])
task(G, 'G04', 'Complete HR notifications expiry reminders and audit APIs', 'Missing producers and queries',
     'HR-007; APR-003, APR-005; AUTH-008; AC-010',
     ['Stage safe outbox events for hiring/change/salary/status/leave/claims/benefits/loan decisions, expiry and policy notices, alongside existing attendance producers.',
      'Implement periodic required/expiring document/certification/insurance detection with approved recipients/window and duplicate prevention. Expose own inbox/read/preferences where permitted.',
      'Add authorized cross-workflow audit search/detail/export for actor/role/scope/IP-device/time/correlation and protected before/after references. Preserve append-only records and sensitive read/export audit.'],
     ['One committed/retried decision yields one intended event; failed transaction sends none. Subjects/lock-screen text exclude salary/CNIC/bank/medical details.',
      'Expiry threshold/replacement/offboarding cases notify the right current authorized recipient. Auditor reads are scoped and cannot mutate audit/decision history.'])
task(G, 'G05', 'Finish API standards security and consumer contract packs', 'Cross cutting implementation and verification',
     'API-001 to API-006; SEC-001 to SEC-012; SRS13',
     ['Publish current OpenAPI and runnable DTO examples for every A-F workflow, including numeric/string enums, cursor versus page responses, evidence previews/fingerprints, versions and idempotent uncertain-result retry.',
      'Finish rate limits for login/OTP/punch/face/file-upload/report paths, standardized problem/code/trace errors, HTTPS/cookie/CSRF/token handling and server-side sensitive field scopes.',
      'Configure least-privilege service identities, network/file/backup encryption and safe logs. Define GPS/face/document/audit retention and approved map-provider privacy; mobile owners consume secure-storage/device permissions as client responsibilities.'],
     ['Correct current contracts work against generated OpenAPI; missing fields, unknown JSON, stale version, conflicting retry key, brute-force and oversized/malicious upload get the documented failure.',
      'IDOR, sensitive export, expired/reused token, file authorization and CSRF tests pass. Logs/errors contain no private payloads/credentials/templates.'])
task(G, 'G06', 'Complete durable jobs operations and restore proof', 'Deployment and resilience',
     'SRS3; NFR-001 to NFR-005; MD-015, MD-016',
     ['Deploy configured notifications/report/recalculation/reminder/face-cleanup workers with durable retry/dead-letter/ownership and observed health. Avoid duplicate jobs on restart or multi-worker races.',
      'Approve employee/site/growth/concurrency/availability targets; measure ordinary API two-second and stable-network punch five-second goals under agreed load, excluding documented large exports.',
      'Run restore drill including SQL, private evidence/files, configurations and required keys; separately follow approved biometric backup/deletion policy. Add startup/readiness checks and observable provider outage handling.'],
     ['Restart/crash/retry does not duplicate notification, consumption or report ownership; queued work resumes and failed work is visible.',
      'An authorized isolated restore reconstructs required source/download history and lawful key access; load report records target, measured timings and failures without calling a build a performance test.'])
task(G, 'G07', 'Execute backend acceptance and publish per part evidence', 'Required completion evidence',
     'SRS16.2; AC-001 to AC-010 including AC-005A and AC-005B',
     ['Reuse existing standalone checks and approved isolated HTTP/SQL/device UAT. Record functional assertions, positive/negative scope, races/replays, persistence and independent restarted-process proof for each part.',
      'Cover onboarding/document/access, office/site/overnight/outstation/exception attendance, live face, corrections, policies, leave/claims/loans, financial closure, own/team views and real exports.',
      'For each task submit endpoint/file changes, reviewed migration/SQL, OpenAPI/examples, actual commands/results, safe evidence and remaining blockers. Do not create permanent test projects without the PM request required by AGENTS.md.'],
     ['Release build exits0 and every required task case has actual passed evidence on the declared source/schema/provider target; blocked/deferred cases remain explicitly open.',
      'The independent reviewer can rerun the evidence and reconcile source versions, SQL rows, private files, exact receipts and consumer contract examples.'])

DEPENDENCIES = {
 'A01':'Current authentication/access services; consumer agreement before changing replay contracts.',
 'A02':'Current file/document services and approved storage/scanner/retention settings.',
 'A03':'A01 current scopes; approved employee-field/category policy decisions.',
 'A04':'A01, A06 and approved new lifecycle transitions; coordinate B05 disposition.',
 'A05':'A02, A06 and approved custody/SIM/certification authority.',
 'A06':'A01 existing access contracts; approved routing/delegation decisions.',
 'A07':'Current office migration/model/SQL; shared migration owner and authorized target.',
 'B01':'A01, A02, A07 and approved office boundary configuration.',
 'B02':'A01, B01 and approved duplicate/offline/anomaly treatment.',
 'B03':'A01, A02, B01, C01; F03 only for monetary locked-period adjustments.',
 'B04':'A01, A02; approved MD-002/004/013 and provider/device deployment.',
 'B05':'A04/B04 lifecycle contracts; approved retention/deletion/hold policy.',
 'B06':'A01, B02; approved MD-003A and accepted zone-scope baseline.',
 'B07':'A01, A02 and existing site-report services; client contract agreement.',
 'B08':'A01, A02, B02; explicit retained/superseded file-import decision.',
 'C01':'A03 taxonomy and existing policy schemas; source contracts can be defined before final producers.',
 'C02':'B01, B02, C01 and D01 exact ordinary leave allocations.',
 'C03':'B01, C01 and approved early/study/make-up instructions.',
 'C04':'B02, C01 and D01 ordinary leave/missing-evidence facts; not D02/D03 benefits.',
 'C05':'B02, B07, C01 and approved pairing/attribution rules; F01 proposed-source contract.',
 'C06':'B01, B02, C01 and approved MD-006/travel/disciplinary examples.',
 'C07':'B02, C01, A03 and approved OT source/category scope; F01 contract for output.',
 'D01':'A02, B01, C01 and existing ordinary leave services; finalized KPI is not a prerequisite.',
 'D02':'D01, C04, A06 and F01; approved six-day grant/election/valuation.',
 'D03':'D01, C04, A06 and F01; approved encashment balance/valuation/consumption.',
 'D04':'D01, C01/C03 and F01; approved recovery/fulfillment/waiver rules.',
 'D05':'D01-D04 as applicable; shared notification and financial/source contracts.',
 'E01':'B02 factual sessions, C01, A03 and approved meal/deployment/cycle rules.',
 'E02':'E01, A02, A06 and F01; G04 notification interface.',
 'E03':'A03, A06, C01 and F01; approved route/provider/commute/valuation policy.',
 'E04':'E03 approved mileage ledger, A06, F01 and workshop/milestone policy.',
 'E05':'A02, A05/A06, F01 and approved custody/incident/liability authority.',
 'E06':'A01, A06 and F01; approved advance/loan/disbursement/repayment terms.',
 'F01':'A01/A02 and current ordinary/late/OT/attestation schemas; agree before final C/D/E producers.',
 'F02':'F01 and accepted C07/D/E source producers for each enabled line type.',
 'F03':'F01, B03 and approved locked-source valuation/target/reversal rules.',
 'F04':'F01 and existing legal/jurisdiction records; formal applicability/activation authority.',
 'F05':'A07, B01-B03, C02, F01-F04 and required D/E inputs; no dependency on a future posting receipt.',
 'F06':'F01-F05 and approved Finance posting/settlement/partial/reversal contract.',
 'G01':'A01/A02 profile sources, D/E/F authoritative own financial producers as applicable.',
 'G02':'A01 current scopes and B01 historical memberships; existing ordinary D01 calendar sources.',
 'G03':'Current export infrastructure; each report depends only on its listed producer, with financial reports completed after C/D/E/F.',
 'G04':'A01/A02/A05 event/expiry sources and current outbox; add each A-F producer as it is accepted.',
 'G05':'A01 and F01 contract agreement; update examples after each module, independent of full client UI.',
 'G06':'Approved target/load/retention/provider deployment; current workers and B05 cleanup interface.',
 'G07':'Required A-F tasks and G01-G06, declared approved target, source/schema/provider evidence.'
}
for package in PARTS:
    for item in package['tasks']:
        item['depends'] = DEPENDENCIES[item['code']]

GAPS = [
 ('Access-assignment Reason omitted by current web client', 'Backend validation exists; client integration blocked', 'A01 / G05'),
 ('SMS owner scanner storage and activation acceptance', 'Existing services; deployment/acceptance incomplete', 'A01 / A02 / G07'),
 ('Controlled fields age expiry bank verification categories', 'Partial policy/operational coverage', 'A03'),
 ('Retirement reactivation rehire archive', 'Transitions/operational paths missing', 'A04'),
 ('Post-hire assets SIM certification renewal and expiry', 'Hiring intake exists; operating lifecycle missing', 'A05 / G04'),
 ('Configurable approval routing and delegation', 'Models only beyond fixed approved flows', 'A06'),
 ('Office roster migration and office leave provenance', 'Now present in source; application/acceptance pending', 'A07 / B01 / D01'),
 ('Office boundary punch correction calculation provenance', 'Site-only consumers remain', 'B01 / B03 / C02'),
 ('Time-tolerance duplicate offline integrity and anomaly review', 'Partial evidence intake', 'B02 / B03'),
 ('Real production face and alternative attendance', 'Development PoC; live provider/device acceptance pending', 'B03 / B04'),
 ('Automatic biometric retirement and physical disposal', 'Use guards exist; durable lifecycle missing', 'B05'),
 ('Green Amber custom zones and full offline zone behavior', 'Operational/Red narrower scope; decision required', 'B06'),
 ('Site report app contracts and scoped queues', 'Backend core exists; complete consumer acceptance pending', 'B07 / G02'),
 ('Attendance file import', 'No operational path; retain or supersede formally', 'B08'),
 ('Early exceptions studying make-up and TDA treatment', 'Missing dedicated producer', 'C03'),
 ('KPI annual late shortage and disciplinary assessments', 'Missing producer/finalized evidence', 'C04 / C06'),
 ('Paired waiting verification attribution and financial effect', 'Model only', 'C05'),
 ('Special cross-midnight resets and normal-paid work travel', 'Continuous evidence exists; special treatment incomplete', 'C06'),
 ('OT broader category/schedule and financial consumption', 'Core formula/review exists; integration incomplete', 'C07 / F02'),
 ('Ordinary leave preview contact reliever team calendar', 'Core leave exists; contract/acceptance extensions', 'D01 / G02'),
 ('Extra six-day or cash benefit', 'Missing lifecycle and balance integration', 'D02'),
 ('Leave encashment recovery and emergency KPI branch', 'Missing specialized services/approved rules', 'D03 / D04 / C04'),
 ('Meal outstation continuous-cycle claims', 'Models only', 'E01 / E02'),
 ('Fuel route commute mileage tuning overhaul custody liability', 'Models only', 'E03 / E04 / E05'),
 ('Advances loans disbursement and repayments', 'Models only', 'E06'),
 ('OT/claims/credits source consumption and exact reconciliation', 'Ordinary/late input core exists; integration incomplete', 'F01 / F02 / F03 / F06'),
 ('Formal legal activation and positive final month close', 'Current readiness gates deliberately block', 'F04 / F05'),
 ('Own profile salary deductions payslip loan data', 'Complete own API producers missing', 'G01'),
 ('All fourteen reports and authorized audit exports', 'Factual export infrastructure exists; full catalogue missing', 'G03 / G04'),
 ('HR notifications document certification insurance reminders', 'Inbox/attendance events exist; HR producers missing', 'G04'),
 ('API standards rate limits privacy load backup and real UAT', 'Implementation/deployment evidence incomplete', 'G05 / G06 / G07'),
 ('Web prepared operations and blank React Native business app', 'Client delivery gap; backend contracts/acceptance dependency', 'G01 / G02 / G05 / G07'),
]

DECISIONS = [
 ('MD-001', 'Onboarding and sensitive-change chain', 'Existing fixed Admin flow retained; decide only new transitions/routes', 'A04 A06'),
 ('MD-002', 'Approved face provider and topology', 'Real production approval/configuration required; Development PoC is not clearance', 'B04'),
 ('MD-003', 'Boundary radius and GPS accuracy', 'Per-site/per-office approved configuration; no global80m mock default', 'B01 B02'),
 ('MD-003A', 'Zone types boundary owner monitoring and exception authority', 'Record Operational/Red deviation or approve missing v1.2 capabilities', 'B06'),
 ('MD-004', 'Liveness match retry enrollment alternative and capture retention', 'Exact approved values/examples required before live activation', 'B03 B04 B05'),
 ('MD-005', 'Correction geofence OT leave allowance authority', 'Preserve HR correction/OT and Admin ordinary leave; define only missing exception/claim authority', 'A06 B03 E02'),
 ('MD-006', 'Continuous duty reset and next-day examples', 'Approve10-hour/4.5-hour/12-extra-hour examples and interactions', 'C06'),
 ('MD-007', 'OT monetary rule', 'Confirmed BASIC/source-calendar-days/9h final whole-rupee rule retained; only new schedule/category/day scope needs decisions', 'C07 F02'),
 ('MD-008', 'Stationed dinner amount and continuous cycle', 'Amount and break/reset definitions remain unconfigured until approved', 'E01'),
 ('MD-009', 'KPI composition annual reset and financial meaning', 'Define base/tier precision/score finalization and whether any effect is monetary', 'C04 D02 D03'),
 ('MD-010', 'Waiting evidence attribution overlap cap and authority', 'Approve factual verification and separate financial/performance treatment', 'C05'),
 ('MD-011', 'Leave benefit cancellation encashment and recovery', 'Ordinary Admin flow retained; define new pool/election/valuation/consumption rules', 'D01 D02 D03 D04'),
 ('MD-012', 'Categories early reporting and studying exceptions', 'Reconcile OFFICE/SITE taxonomy and approved category eligibility', 'A03 C03 C07 E03'),
 ('MD-013', 'Face roles offboarding retirement and deletion', 'Approve authority/use/retention/deletion/backup treatment', 'A04 B04 B05'),
 ('MD-014', 'Finance handoff authority partial reversal and legal activation', 'Preserve exact permissioned receipt principles; formal applicability and new effects need authority', 'F01 F04 F05 F06'),
 ('MD-015', 'Retention for GPS punches evidence documents finance audit', 'Approve durations/holds/disposal/restore before cleanup activation', 'A02 B05 B06 G04 G06'),
 ('MD-016', 'Production size growth peak traffic and availability', 'Approve measurable targets and acceptance load', 'G06'),
]
REPORTS = [
 ('01','Daily attendance','B/C/D','Present/late/absence/leave/exception, worked/payable, source and scoped filters'),
 ('02','Monthly register','B/C/D/F','Employee-day matrix, status totals, worked/payable time and exact lock state'),
 ('03','Punch audit','B','Raw IN/OUT, device/source, authorized GPS and correction trail'),
 ('04','Operational zones','B06','Assignments/events/warnings/acknowledgements/exceptions and reviewer decisions'),
 ('05','Site reporting','B07','Arrival/departure/delay/access/incident, sync state and review history'),
 ('06','Late and KPI','C04/C06','Policy-year counts, tier calculation and approved disciplinary status'),
 ('07','Overtime','C07/F02','Requested/approved/rejected/payable awards by employee/site/job'),
 ('08','Meal allowance','E01/E02/F','Eligibility cycle, calculation, approval and actual payment status'),
 ('09','Leave','D','Balance/utilization/evidence/absence and benefit/encashment eligibility'),
 ('10','Employee master','A/G01','Only authorized current/profile fields; sensitive fields excluded by default'),
 ('11','Document expiry','A02/A05/G04','Required identity/certificate/insurance expiry and reminder state'),
 ('12','Fuel and mileage','E03/E04/E05','Approved route/km/liters, maintenance milestones and adjustments'),
 ('13','Payroll handoff','F','Locked exact inputs, OT/allowance/deductions and source/receipt trace IDs'),
 ('14','Audit','A/B/G04','Sensitive changes, policies, corrections, approvals, exports and access'),
]
ACS = [
 ('AC-001','A01 A02 A03 A06 A07','Complete draft -> verified documents -> independent approval -> corrected access assignment -> activation; no data loss'),
 ('AC-002','B01 B02 B04 B05 G05','Actual live one-to-one/device Check In and server event; Check Out without face; own/scoped event visibility'),
 ('AC-003','B03 F03','Correction changes interpreted revisions and retains original raw evidence, reviewer and adjustment trail'),
 ('AC-004','B01 C01-C07 D E01 F','Approved office/site ordinary/late/half-day/hourly/OT/allowance examples use exact effective versions'),
 ('AC-005','B01 B02 B03','Outside/poor-location evidence has clear review status; unauthorized review and alternate-process bypass fail'),
 ('AC-005A','B06','Approved zone scope or signed deviation; separate event/notice with safe own instruction and unchanged punch'),
 ('AC-005B','B07 G05','All arrival/departure/delay/incident reports have API, evidence and reviewed history without external messaging'),
 ('AC-006','D E06 F G01 G02','Office/site employees see only own attendance/finance/leave/OT/advances/loans; ID tampering denied'),
 ('AC-007','F01-F06','Accounts consumes exact locked inputs/receipts while raw punch/policy writes stay denied'),
 ('AC-008','G03 G05 G06','Actual14 report types generate usable authorized files; forbidden columns/downloads denied'),
 ('AC-009','A04 B02 B04','Effective suspension/termination denies new/replayed attendance and preserves authorized history'),
 ('AC-010','A06 B C D E F G04','Sensitive reads/changes/policies/decisions/exports trace actor/source/version/before-after and immutable audit'),
]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.72), Inches(.72)
sec.header_distance, sec.footer_distance = Inches(.25), Inches(.25)
styles = doc.styles
for style in styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s = styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0,0,0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),'Calibri')
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),'Calibri')
styles['Normal'].font.size=Pt(11)
styles['Normal'].paragraph_format.space_after=Pt(5)
styles['Normal'].paragraph_format.line_spacing=1.06
styles['Title'].font.size=Pt(25)
styles['Title'].paragraph_format.space_after=Pt(12)
styles['Subtitle'].font.size=Pt(13)
styles['Subtitle'].paragraph_format.space_after=Pt(10)
for name,size in [('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
    styles[name].font.size=Pt(size)
    styles[name].font.bold=True
    styles[name].paragraph_format.space_before=Pt(11)
    styles[name].paragraph_format.space_after=Pt(5)
    styles[name].paragraph_format.keep_with_next=True
styles['Normal'].paragraph_format.widow_control=True

def clean_text(text):
    text = str(text)
    if text.startswith(('src/','docs/','AGENTS.md')):
        return text
    protected = []
    pattern = r'\b(?:[A-G]\d{2}(?:\.T?\d+)?|(?:MD|AC|API|OFF|SITE|HR|AUTH|ROLE|REP|NFR|LEAVE|VEH|ALLOW|ZONE|FACE|ATT|MOB|SEC)-\d+[A-Z]?|v\d+(?:\.\d+)+)\b|\b[\w./-]+\.(?:cs|md|json|sql|docx)\b'
    def protect(match):
        protected.append(match.group(0))
        return chr(0xE000+len(protected)-1)
    text = re.sub(pattern,protect,text)
    for joined,separated in [('minimumfour-month','minimum four-month'),
                             ('exactlyfour-month','exactly four-month'),
                             ('approvedkm','approved km'),('outstationkm','outstation km')]:
        text=text.replace(joined,separated)
    text=re.sub(r'(?<=[a-z])(?=PKR(?:\b|\d))',' ',text)
    text = re.sub(r'(?<=[A-Za-z])(?=\d)',' ',text)
    text = re.sub(r'(?<=\d)(?=[A-Za-z])',' ',text)
    for i,original in enumerate(protected):
        text=text.replace(chr(0xE000+i),original)
    return text

def para(text, boldlead=None, small=False):
    text=clean_text(text)
    p=doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold=True
        p.add_run(text[len(boldlead):])
    else:p.add_run(text)
    if small:
        for run in p.runs: run.font.size=Pt(9.5)
        p.paragraph_format.line_spacing=1
    if text in ['Acceptance','Existing touchpoints']:
        p.paragraph_format.keep_with_next=True
    return p
def heading(text,level=1):return doc.add_heading(clean_text(text),level)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=clean_text(v)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcp.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            tcp.append(margins)
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'20384F' if ri==0 else ('F1F4F7' if ri%2==0 else 'FFFFFF'));tcp.append(shading)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.02
                for run in p.runs:
                    run.font.size=Pt(9.5)
                    if ri==0:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

# Page-number fields aid developer task references in a long handoff.
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('Cosmix HR Backend Plan  |  Page ');run.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Cosmix HR Backend Completion Plan'
doc.core_properties.subject='Implementation Parts A through G and developer acceptance'
doc.core_properties.author='Cosmix Engineering'
doc.core_properties.keywords='HR Attendance Backend SRS Implementation A B C D E F G'

doc.add_paragraph('Cosmix HR Backend\nCompletion Plan',style='Title')
doc.add_paragraph('Implementation Parts A through G',style='Subtitle')
para('Developer handoff  |  13 October 2026  |  Requirement baseline SRS version 1.2',small=True)
para('Goal: Complete the remaining HR and attendance backend services, data persistence, approved business rules, client contracts and operational acceptance so the developers can deliver independently reviewable packages.',boldlead='Goal:')
para('PM note: Is document ke Parts A se G developers ko alag assign kiye ja sakte hain. Har task ki numbered lines mein required change, dependency aur acceptance diya gaya hai. Jo service pehle se bani hui hai usko reuse karein; pending rule ya provider ko approved evidence ke baghair complete mark na karein.')
para('The largest remaining backend work is office attendance provenance, special policy/leave engines, allowance/vehicle/advance/loan services, full financial consumption and month closure, own financial/profile APIs, all fourteen reports and operational acceptance. Employee administration and the bounded attendance core already provide reusable implementations.')
heading('How to use this handoff',2)
para('Assign task IDs such as B01 or E03 to an owner. Each implementation line is separately trackable. Finish the stated acceptance cases before marking a task accepted. Preserve current approved behavior and update the status ledger with the exact source/schema/provider evidence.')
table(['Part','Deliverable','Primary owner'],[(p['letter'],p['title'],p['owner']) for p in PARTS],[.45,3.45,2.93])
para('Scope: Backend source and deployed API acceptance. Full React/React Native screens are separate client work. This plan includes their required contracts and client-dependent acceptance; backend completion does not declare the whole platform or mobile app complete.',boldlead='Scope:')

doc.add_page_break()
heading('Current baseline and completion definition')
para('Primary implementation root: C:/Users/Huma.Khan/Downloads/Cosmix_Erp_software/cosmix-backend. Use this latest local source rather than the older separate Downloads/cosmix-backend copy. Technical stack is .NET10, C#, SQL Server and EF Core; business logic stays in Domain/Application with module-specific repositories and API composition.')
para('Existing delivery: authentication/access/recovery; masters/holidays/company configuration; hiring/documents/current employee changes/compensation/limited status/site assignment; shifts/site rosters/site geofences/new office roster source; immutable punches/sessions/factual summaries; corrections/OT/Operational-Red zones/site reports; direct policies/ordinary leave; saved ordinary-late inputs, no-charge handoff/manual Finance attestations; factual export jobs and own notification inbox.',boldlead='Existing delivery:')
para('Latest delta: AddOfficeRosterHistory migration20261013101327, its snapshot/SQL and office leave-source retention now exist. Earlier audit findings that these were absent are superseded. Review and accept them; do not recreate them. Office capture/calculation/manual-correction paths still depend on site evidence and need B01/B03.',boldlead='Latest delta:')
para('Recorded bounded acceptance: the latest repository verification record reports a synthetic development target upgraded through migration21, actual password/MFA/CSRF/logout checks and unassigned-shift retry/concurrency checks. It also records a later independent evidence check blocked by office model/schema coordination. These outcomes do not establish complete office, face, claims, financial or whole-HR acceptance. Verify current target/schema state before executing any new acceptance.')
table(['Completion level','Required evidence'],[
 ('Source implemented','Real validated service, controller, repository, DI, migration/model agreement, audit, idempotency/concurrency and current OpenAPI/contracts. Folders/entities alone are insufficient.'),
 ('Deployed backend accepted','Reviewed authorized schema/provider/configuration; real authenticated persistence, permissions, concurrency, files/jobs, lifecycle and financial readiness cases pass.'),
 ('Whole platform accepted','Web/mobile owners complete the consuming screens/device flows; full SRS ACs and approved scope deviations have end-to-end evidence. This is a separate status.')],[1.65,5.18])
heading('Global implementation constraints',2)
for i,text in enumerate([
 'Keep one employee master; use stable EmployeeId in attendance, finance and mobile. Keep sensitive HR fields in separately permissioned projections and private encrypted storage.',
 'Raw punches, decisions and retained evidence are append-only. Effective-dated assignments/policies/compensation must not silently rewrite prior periods.',
 'Preserve confirmed direct HR scheduling/policy publication, Admin ordinary leave, HR correction/OT authority, office OT ineligibility and no finalized-month reopening.',
 'Preserve confirmed ordinary BASIC/actual-month-days precision and larger-only ordinary/late overlap. Preserve OT BASIC/source-month-calendar-days/9h exact-time formula with final whole-rupee rounding.',
 'Retain valid no-charge handoff. Saved calculation, eligibility, approval, handoff or export success does not mean posting, settlement or bank payment.',
 'Provider/authority/required source unavailable must remain explicit and fail closed. Do not use synthetic provider success or invented defaults to release pay or accept Check In.',
 'Coordinate Program/DI/DbContext/permissions/migrations through one shared-file owner. Never rewrite an applied migration or erase retained data to obtain a passing check.',
 'Use existing standalone checks and approved acceptance scripts. AGENTS.md prohibits recreating permanent test projects without a PM request; SRS calculation/integration/contract/security/UAT cases still need executable evidence.',
 'No automatic bank disbursement, attendance-machine integration, external-message attendance channel, unrestricted tracking or automated termination/medical diagnosis is part of this completion work.'
],1):para(f'R{i:02d}. {text}')
heading('Review focus',2)
for i,text in enumerate([
 'Lost response and later revoked permission: retry/reconciliation cannot duplicate a grant, punch, claim or financial consumption (A01, B02, F01/F06).',
 'Office/site transfer during overnight duty or before leave approval: exact retained membership/shift/boundary evidence must reject stale or conflicting interpretation (B01, C02, D01).',
 'Later policy/KPI/benefit revision: a displayed score, new salary or pending election cannot rewrite a retained source or duplicate benefit/encashment (C04, D02/D03, F03).',
 'Scanner, face provider or encrypted storage unavailable: no clean document or successful Check In is invented; alternative approval has its own authorized evidence (A02, B03-B05).',
 'Partial/reversed Finance receipt and revoked export scope: only exact confirmed source lines are settled and sensitive downloads remain denied (F01-F06, G03/G05).'
],1):para(f'V{i:02d}. {text}')

heading('Gap register and task ownership')
para('The register below maps every major gap from the audit to an implementation task. Code gaps, policy decisions, client dependencies and deployment acceptance are distinct work types; resolve each using its stated evidence.')
table(['Gap','Current state','Task IDs'],GAPS,[2.75,2.98,1.10])

heading('Management decisions and baseline deviations')
para('Management supplies applicability, authority and expected policy examples; developers implement the approved records. Preserve confirmations listed above. A decision can close a gap through supported implementation or an explicitly approved SRS deviation with replacement acceptance. An unapproved omission cannot be marked complete.')
table(['ID','Decision area','Required resolution','Tasks'],DECISIONS,[.63,1.56,3.43,1.21])
heading('Additional decisions for new services',2)
para('Decide advance/loan terms and disbursement/repayment/amendment authority; meal cycle/home-break closure; fuel-price/value and commute adjustment rules; workshop/mileage milestone handling; custody/liability review; extra-benefit and encashment valuation/pool consumption; recovery/make-up credit/deadline/waiver; and formal legal activation/partial-posting/reversal authority. Retain the explicit SRS amounts/thresholds; do not invent missing values.')
para('Recorded SRS differences: office-wise plus site scheduling replaced historical individual precedence; Operational/Red currently narrows Green/Amber/custom-zone scope; direct HR policy publication and fixed ordinary approval flows differ from a generic approval design; finalized months do not reopen. Obtain a traceable accepted baseline for these differences before claiming literal v1.2 conformance.')

for p in PARTS:
    doc.add_page_break()
    heading('Part '+p['letter']+' '+p['title'])
    para('Owner: '+p['owner'],boldlead='Owner:')
    para('Dependencies: '+p['prerequisite'],boldlead='Dependencies:')
    para('Current state: '+p['current'],boldlead='Current state:')
    para('Existing touchpoints',boldlead='Existing touchpoints')
    for anchor in p['anchors']:para(anchor,small=True)
    for t in p['tasks']:
        task_start=len(doc.paragraphs)
        heading(t['code']+' '+t['title'],2)
        para('Work type: '+t['kind']+'  |  SRS: '+t['srs'],small=True).paragraph_format.keep_with_next=True
        para('Task dependencies: '+t['depends'],small=True).paragraph_format.keep_with_next=True
        for i,w in enumerate(t['work'],1):para(f"{t['code']}.{i} {w}")
        para('Acceptance',boldlead='Acceptance')
        for i,c in enumerate(t['checks'],1):para(f"{t['code']}.T{i} {c}")
        for task_paragraph in doc.paragraphs[task_start:-1]:
            task_paragraph.paragraph_format.keep_with_next=True
        doc.paragraphs[-1].paragraph_format.keep_with_next=False
    if p['letter']=='G':
        heading('Fourteen report delivery checklist',2)
        table(['No','Report','Sources','Required content'],REPORTS,[.38,1.43,1.05,3.97])

doc.add_page_break()
heading('Developer assignment and integration order')
para('Use the existing two-developer ownership model. Developer1 owns Authentication/Application HR and related services; Developer2 owns Attendance/provider/jobs. One nominated owner controls each shared-file/migration window. A module owner supplies both its implementation and safe reproducible acceptance evidence.')
table(['Wave','Developer 1','Developer 2','Review gate'],[
 ('1','A01-A06 contracts/remaining HR; draft E/F source interfaces','A07 schema review with migration owner; B01-B03 office evidence','Authentication/storage/schema/source contracts agree'),
 ('2','C01 configuration; D01 ordinary leave; E factual intake/eligibility','B04-B08 provider/retention/zones/reports; C02-C07 calculations','Approved decisions plus factual/policy source cases'),
 ('3','D02-D05 benefits; E01-E06 claims/loans; F source adapters','C finalized outputs; F02-F05 consumption/adjustment/lock readers','Exact sources, no duplicate consumption and no lock cycle'),
 ('4','F06 financial views/receipts; G01/G02 HR ESS/team/report producers','G03/G04 attendance reports/notifications; G05/G06 operations','All required producers/permissions/jobs integrated'),
 ('5','G07 HR/claims/finance/own-data acceptance','G07 attendance/provider/lock/export/restore acceptance','Independent review and declared target evidence')],[.45,2.10,2.10,2.18])
para('Parallel work: G05 contract documentation/security queries, G03 report framework, G04 outbox/templates and G06 operational preparation can start early. Financial/KPI/meal/loan reports cannot be accepted until their producers exist. Agree F01 before C/D/E teams generate financial adapters. Do not assign both developers concurrent edits to the same migration snapshot or permission registry.')
heading('Task status and handoff requirements',2)
para('Allowed status: Not started; In implementation; Source verified; Awaiting decision/provider/schema; Backend accepted; Consumer acceptance pending. Record the owner and evidence link per task. A package is accepted only when all of its required task outcomes are accepted or covered by a signed scope replacement.')
for i,text in enumerate([
 'List completed task IDs and exact changed services/endpoints. Include source revision and declared current migration/schema/provider versions.',
 'Provide reviewed incremental migration and rollback/retention implications; record which target actually applied it. Preserve prior applied sources.',
 'Publish current OpenAPI, DTO/enum/page-or-cursor examples, required permissions/scopes, evidence/preview/version fields and replay/conflict examples.',
 'Record executed checks and meaningful results: positive flow, permission denial, stale/race/idempotency, durable storage, source trace and restart acceptance. Never relabel a prepared test as passed.',
 'List unresolved management values/providers, client dependencies and blocked cases with their owning task/decision IDs. Do not report a package complete while required cases remain blocked.'
],1):para(f'H{i:02d}. {text}')

doc.add_page_break()
heading('Acceptance mapping and final backend signoff')
para('Use the following mapping for acceptance records. Backend-only signoff verifies the server contracts and effects; actual app/camera/UI outcomes need the named client/device owner as an additional acceptance step. All signed baseline deviations travel with the result.')
table(['SRS AC','Tasks','Required proof'],ACS,[.71,1.61,4.51])
heading('Final backend completion checklist',2)
for i,text in enumerate([
 'Every gap-register row has an accepted task or formally approved scope replacement. All required policy values and authority records are resolved; unavailable modules remain open, not counted as zero liability.',
 'Runtime EF model, snapshot, reviewed migration range and target history agree. Existing data/source bytes are preserved and independent restarted-process persistence/receipt evidence is accepted.',
 'Office/site schedules, accepted punches, corrections, policy outputs, leave/claims/loan and financial source consumption work through real registered services with scope/version/retry/trace enforcement.',
 'Approved real face provider, device/liveness/security cases and retention/disposal operations are accepted. If biometric acceptance is formally deferred, backend production attendance acceptance remains explicitly incomplete.',
 'A positive complete nonempty month finalizes with exact immutable input/authority evidence; Accounts can consume locked records and post/settle only exact authorized receipts. No dependency cycle or reopening is required.',
 'All fourteen real reports, own/team API projections, safe HR notifications/reminders and protected audit queries produce authoritative scoped results; no UI-only/mock success is used.',
 'Release build, standalone calculation checks, real authenticated SQL/API/security/concurrency/files/jobs and approved load/restore/UAT cases pass on the declared target. Web/mobile consumer work has a separate acceptance status.',
 'PM/HR, technical reviewer and the authorized Finance/legal/deployment owners sign their applicable acceptance records. A document, source build or table count alone cannot close the backend.'
],1):para(f'S{i:02d}. {text}')

doc.add_page_break()
heading('Requirement traceability and source index')
table(['Requirement family','Implementation ownership'],[
 ('ROLE / AUTH','A01 A06 G01 G02 G05 G07'),('HR / employee fields / masters','A02-A07 G01 G04'),
 ('SHIFT / sites / raw attendance / corrections','A07 B01-B03 B08 C02 F03 F05'),
 ('FACE / biometric security','A04 B03-B05 G05-G07'),('ZONE / site reports','B06 B07 G02 G04 G05'),
 ('OFF / SITE policy rules / OT','B01 C01-C07 D E01 F01-F03'),
 ('LEAVE','C04 D01-D05 F02 F03 G01 G03 G04'),('ALLOW / VEH / advances loans','E01-E06 F01-F03 F06 G01 G03'),
 ('APR / audit','A06 all module decisions G04 G05'),('API / data model / constraints','A01 A07 all module contracts G05 G07'),
 ('REP / fourteen report catalogue','G03 plus authoritative A-F producers'),('SEC / NFR / testing','A01 A02 B04-B06 G05-G07'),
 ('MOB server-facing requirements','B01-B07 D E F G01 G02 G05; actual client UI/device implementation tracked separately'),
 ('AC / MD','Acceptance and decision registers in this document')],[2.04,4.79])
para('Requirement source: HR_Attendance_Employee_Mobile_SRS_v1.2.docx supplied by the PM, dated1 October2026. Sections5-10 define employee/attendance/policies/leave/claims; sections11-17 define consumer/API/data/report/security/testing/acceptance; section19 contains management decisions. Signed source policy references and approved later scope records must accompany implementation.')
para('Repository evidence and approved ownership records',boldlead='Repository evidence and approved ownership records')
for path in [
 'AGENTS.md; docs/Architecture.md; docs/SRS-Module-Map.md; docs/Work-Ownership.md',
 'docs/HR-Gap-Closure-2026-10-13.md',
 'docs/superpowers/specs/2026-10-13-hr-attendance-setup-design.md',
 'docs/attendance-completion/Verification.md; SRS-Remaining-Decisions-20261012.md',
 'docs/attendance-capture/Face-PoC-Completion-20261013.md; Face-Storage-Lifecycle.md',
 'src/Cosmix.Infrastructure/Database/Migrations/20261013101327_AddOfficeRosterHistory.cs',
 'src/Cosmix.Application/HR/Leaves/OfficeLeaveScheduleSource.cs',
 'src/Cosmix.Application/Attendance/PayrollInputs/AttendanceMonthLockFinancialReader.cs',
 'docs/attendance-completion/Payroll-Attestations-Implementation.md; Payroll-Reconciliation-Implementation.md',
 'docs/Employee-Hiring.md; Employee-Updates.md; Employee-Compensation.md; Organization-Masters.md',
 'src/Cosmix.Application/HR/{Allowances,VehicleFuel,AdvancesLoans,EmployeeSelfService,Reports}/',
 'src/Cosmix.Application/Attendance/{Reports,Corrections,Zones,SiteReports,Overtime,MonthLocks}/'
]:para(path,small=True)
para('Backend percentages from the earlier audit are scope estimates, not work-hour budgets or acceptance certificates. This handoff uses current task outcomes and source/deployment status so developers can track completion without rebuilding already-delivered features.')

doc.save(OUT)
data = dict(parts=PARTS,gaps=GAPS,decisions=DECISIONS,reports=REPORTS,acceptance=ACS)
(Path(__file__).parent/'plan-content.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUT),'parts':len(PARTS),'tasks':sum(len(p['tasks']) for p in PARTS),
                  'gaps':len(GAPS),'decisions':len(DECISIONS),'reports':len(REPORTS),
                  'acceptance_criteria':len(ACS),'bytes':OUT.stat().st_size}))
